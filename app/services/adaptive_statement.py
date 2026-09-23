"""Adaptive reader for bank statements that do not match a known template.

The reader rebuilds a table from an Excel sheet or from the word positions in a
PDF, then decides which numeric columns are movements and which are a running
balance. When the statement prints an opening and a closing balance, the
decision is the one that makes those two figures agree.
"""

from __future__ import annotations

import re
import unicodedata
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from io import BytesIO
from pathlib import Path

import fitz
from openpyxl import load_workbook

from app.schemas.statement import ParsedStatement, StatementMetadata, StatementTotals, StatementTransaction
from app.services.document_service import DocumentParseError
from app.services.halyk_fiz_statement_service import looks_like_halyk_personal
from app.services.legal_statement import LEGAL, holder_kind, legal_operation, legal_totals

_EXCEL_EPOCH = datetime(1899, 12, 30)
_MAX_EXCEL_ROWS = 5000
_MAX_PDF_PAGES = 40
_HEADER_SCAN_ROWS = 45

# In a client statement дебет is money leaving the account and кредит is money
# coming in. A running balance is not turnover and must not be added to either.
_HEADER_PHRASES: tuple[tuple[str, str], ...] = (
    ("income", "сумма поступления"),
    ("expense", "сумма списания"),
    ("income", "сумма кредита"),
    ("expense", "сумма дебета"),
    ("amount", "сумма в валюте счета"),
    ("amount", "сумма операции"),
    ("amount", "сумма платежа"),
    ("amount", "сумма транзакции"),
    ("date", "дата операции"),
    ("date", "дата проводки"),
    ("date", "дата документа"),
    ("date", "дата транзакции"),
    ("value_date", "дата валютирования"),
    ("value_date", "дата обработки"),
    ("detail", "назначение платежа"),
    ("detail", "основание платежа"),
    ("detail", "төлем мақсаты"),
    ("counterparty", "наименование бенефициара"),
    ("counterparty", "наименование контрагента"),
    ("counterparty", "наименование отправителя"),
    ("counterparty", "контрагент"),
    ("counterparty", "бенефициар"),
    ("counterparty", "отправитель"),
    ("counterparty", "получатель"),
    ("document", "номер документа"),
    ("balance", "остаток после операции"),
    ("balance", "исходящий остаток"),
    ("balance", "входящий остаток"),
    ("operation", "вид операции"),
    ("operation", "тип операции"),
    ("currency", "валюта операции"),
    ("income", "поступление"),
    ("income", "зачисление"),
    ("income", "приход"),
    ("income", "кредит"),
    ("income", "credit"),
    ("income", "inflow"),
    ("income", "кіріс"),
    ("expense", "списание"),
    ("expense", "расход"),
    ("expense", "дебет"),
    ("expense", "debit"),
    ("expense", "outflow"),
    ("expense", "шығыс"),
    ("amount", "сумма"),
    ("amount", "amount"),
    ("balance", "остаток"),
    ("balance", "balance"),
    ("detail", "назначение"),
    ("detail", "описание"),
    ("detail", "детали"),
    ("detail", "description"),
    ("detail", "details"),
    ("detail", "комментарий"),
    ("detail", "основание"),
    ("counterparty", "counterparty"),
    ("counterparty", "merchant"),
    ("operation", "операция"),
    ("operation", "operation"),
    ("currency", "валюта"),
    ("currency", "currency"),
    ("date", "күні"),
    ("date", "дата"),
    ("date", "date"),
    ("document", "документ"),
    ("document", "document"),
)
_PHRASES = tuple(sorted(_HEADER_PHRASES, key=lambda item: len(item[1]), reverse=True))

_TABLE_END_WORDS = (
    "обороты",
    "итого",
    "всего",
    "исходящий остаток",
    "остаток на конец",
    "исходящее сальдо",
)

_NOT_A_HEADER = (
    "последнего движения",
    "дата выписки",
    "дата формирования",
    "дата печати",
    "дата создания",
)

_OPENING_LABELS = (
    "входящий остаток",
    "входящее сальдо",
    "остаток на начало",
    "начальный остаток",
    "opening balance",
    "баланс на начало",
)
_CLOSING_LABELS = (
    "исходящий остаток",
    "исходящее сальдо",
    "остаток на конец",
    "конечный остаток",
    "closing balance",
    "баланс на конец",
)
_HOLDER_LABELS = (
    "владелец счета",
    "владелец счёта",
    "наименование клиента",
    "account holder",
    "клиент",
    "фио",
)
_ACCOUNT_LABELS = (
    "номер счета",
    "номер счёта",
    "текущий счет",
    "текущий счёт",
    "лицевой счет",
    "лицевой счёт",
    "счет",
    "счёт",
)

_BANKS: tuple[tuple[str, str], ...] = (
    ("forte", "Forte Bank"),
    ("jusan", "Jusan Bank"),
    ("freedom", "Freedom Bank"),
    ("halyk", "Halyk Bank"),
    ("народный банк", "Halyk Bank"),
    ("kaspi", "Kaspi"),
    ("центркредит", "Банк ЦентрКредит"),
    ("centercredit", "Банк ЦентрКредит"),
    ("home credit", "Home Credit Bank"),
    ("хоум кредит", "Home Credit Bank"),
    ("евразий", "Евразийский банк"),
    ("bereke", "Bereke Bank"),
    ("altyn", "Altyn Bank"),
    ("отбасы", "Отбасы банк"),
    ("bank rbk", "Bank RBK"),
    ("rbk bank", "Bank RBK"),
)

_CURRENCY_RE = re.compile(r"kzt|usd|eur|rub|gbp|cny|тенге|тг|₸|\$|€", re.IGNORECASE)
_DATE_FULL_RE = re.compile(r"^(\d{2})[./-](\d{2})[./-](\d{2,4})(?:\s+\d{2}:\d{2}(?::\d{2})?)?$")
_DATE_ISO_RE = re.compile(r"^(\d{4})-(\d{2})-(\d{2})")
_DATE_IN_TEXT_RE = re.compile(r"\d{2}[./]\d{2}[./]\d{2,4}")
_PERIOD_FROM_TO_RE = re.compile(
    r"с\s+(\d{2}[./]\d{2}[./]\d{2,4})\s+по\s+(\d{2}[./]\d{2}[./]\d{2,4})",
    re.IGNORECASE,
)
_PERIOD_DASH_RE = re.compile(
    r"(?:период|period)\s*[:\s]\s*(\d{2}[./]\d{2}[./]\d{2,4})\s*[-–—]\s*(\d{2}[./]\d{2}[./]\d{2,4})",
    re.IGNORECASE,
)
_IBAN_RE = re.compile(r"\bKZ[0-9A-Z]{18}\b", re.IGNORECASE)
_TAX_ID_LABEL_RE = re.compile(
    r"^(?:иин|бин|инн|iin|bin)(?:\s*/\s*(?:иин|бин|iin|bin))?"
    r"(?:\s+(?:клиента|владельца|организации))?\s*:?\s*(.*)$",
    re.IGNORECASE,
)
_TAX_ID_VALUE_RE = re.compile(r"(?<!\d)(\d{12})(?!\d)")
_AMOUNT_TOKEN_RE = re.compile(
    r"(?<!\d)(?P<neg>\(|[+-])?\s*(?P<num>\d{1,3}(?:[ \u00a0]\d{3})+(?:[.,]\d{2})?|\d+[.,]\d{2})\)?(?!\d)"
)
_LINE_DATE_RE = re.compile(
    r"^(\d{2}[./]\d{2}[./]\d{2,4})(?:\s+\d{2}:\d{2}(?::\d{2})?)?\b\s*(.*)$"
)

_GENERIC_OPERATION_WORDS = frozenset({
    "операция",
    "operation",
    "type",
    "тип",
    "вид операции",
    "тип операции",
})


@dataclass
class MoneySlot:
    index: int
    role: str
    label: str


@dataclass
class Layout:
    date_col: int
    money: list[MoneySlot]
    detail_cols: list[int] = field(default_factory=list)
    operation_col: int | None = None
    document_col: int | None = None
    value_date_col: int | None = None
    counterparty_col: int | None = None
    purpose_col: int | None = None
    headerless: bool = False


@dataclass
class RawRow:
    date: str
    values: dict[int, float | None]
    detail: str
    operation_text: str
    document_number: str | None
    processing_date: str | None
    # Контрагент и назначение платежа по отдельности, если у них свои колонки:
    # вид для юрлица кладёт их в разные колонки.
    counterparty: str = ""
    purpose: str = ""


@dataclass
class Plan:
    mode: str
    priority: int
    transactions: list[StatementTransaction]
    residual: float | None
    pseudo: float
    ignores_balance: bool
    all_positive: bool
    income_label: str
    expense_label: str
    amount_label: str


@dataclass
class ReadResult:
    statement: ParsedStatement
    headerless: bool
    reconciled: bool | None


def detect_adaptive_statement(filename: str, content: bytes) -> float:
    extension = Path(filename).suffix.lower()
    if extension not in {".pdf", ".xlsx", ".xlsm"}:
        return 0.0
    try:
        if _specialist_fingerprint(filename, content):
            return 0.2
        result = _read(filename, content)
    except Exception:
        return 0.0
    if result is None:
        return 0.0
    if result.headerless and not result.reconciled:
        return 0.52
    if result.reconciled:
        return 0.64
    return 0.62


def parse_adaptive_statement(filename: str, content: bytes) -> ParsedStatement:
    result = _read(filename, content)
    if result is None:
        raise DocumentParseError(
            "Не удалось найти в файле таблицу операций. "
            "Нужна колонка с датами и колонка с суммой, дебетом или кредитом."
        )
    return result.statement


def _read(filename: str, content: bytes) -> ReadResult | None:
    extension = Path(filename).suffix.lower()
    if extension in {".xlsx", ".xlsm"}:
        return _read_excel(filename, content)
    if extension == ".pdf":
        return _read_pdf(filename, content)
    return None


def _read_excel(filename: str, content: bytes) -> ReadResult | None:
    best: ReadResult | None = None
    best_key: tuple[int, int] | None = None
    for rows in _excel_sheets(content):
        result = _interpret_grid(filename, rows)
        if result is None:
            continue
        key = (
            1 if result.reconciled else 0,
            result.statement.metadata.transaction_count,
        )
        if best_key is None or key > best_key:
            best = result
            best_key = key
    return best


def _read_pdf(filename: str, content: bytes) -> ReadResult | None:
    positional = _pdf_grid(content)
    positioned = _interpret_grid(filename, positional[0], extra_text=positional[1]) if positional else None
    linear = _read_pdf_lines(filename, content)
    if positioned is None:
        return linear
    if linear is None:
        return positioned
    positioned_count = positioned.statement.metadata.transaction_count
    linear_count = linear.statement.metadata.transaction_count
    if positioned.headerless and not positioned.reconciled and linear_count > positioned_count:
        return linear
    return positioned


def _excel_sheets(content: bytes) -> list[list[list[object]]]:
    workbook = load_workbook(BytesIO(content), data_only=True, read_only=True)
    sheets: list[list[list[object]]] = []
    try:
        for sheet in workbook.worksheets[:6]:
            rows: list[list[object]] = []
            for index, row in enumerate(sheet.iter_rows(values_only=True)):
                rows.append(list(row))
                if index + 1 >= _MAX_EXCEL_ROWS:
                    break
            if any(any(cell not in (None, "") for cell in row) for row in rows):
                sheets.append(rows)
    finally:
        workbook.close()
    return sheets


def _interpret_grid(
    filename: str,
    rows: list[list[object]],
    extra_text: str = "",
) -> ReadResult | None:
    if not rows:
        return None
    header_index, roles, extras = _find_header(rows)
    if header_index is None:
        layout = _infer_layout(rows)
        body = rows
        preamble = rows[:8]
        headerless = True
    else:
        layout = _layout_from_roles(rows[header_index], roles, extras)
        body = rows[header_index + 1 :]
        preamble = rows[:header_index]
        headerless = False
    if layout is None:
        return None

    merged = _merge_continuations(body, layout)
    raw_rows = [_raw_row(row, layout) for row in merged]
    raw_rows = [row for row in raw_rows if row is not None]
    if not raw_rows:
        return None

    preamble_text = _rows_text(preamble)
    context_text = "\n".join(part for part in (preamble_text, extra_text) if part)
    opening = _labeled_amount(preamble, _OPENING_LABELS)
    closing = _labeled_amount(preamble, _CLOSING_LABELS)
    if opening is None:
        opening = _amount_after_label(context_text, _OPENING_LABELS)
    if closing is None:
        closing = _amount_after_label(context_text, _CLOSING_LABELS)

    plan = _choose_plan(raw_rows, layout, opening, closing, headerless)
    if plan is None or not plan.transactions:
        return None

    period_start, period_end = _period(context_text)
    if period_start is None or period_end is None:
        dated = [_parse_stamp(row.date) for row in raw_rows]
        dated = [item for item in dated if item is not None]
        if dated:
            period_start = period_start or min(dated).strftime("%d.%m.%Y")
            period_end = period_end or max(dated).strftime("%d.%m.%Y")

    currency = _currency(context_text)
    holder = _labeled_text(preamble, _HOLDER_LABELS) or _text_after_label(context_text, _HOLDER_LABELS)
    account = _iban(context_text) or _labeled_text(preamble, _ACCOUNT_LABELS)
    bank = _bank_name(f"{filename}\n{preamble_text}")
    title = f"Выписка {bank}" if bank else "Банковская выписка"
    note = _compose_note(plan, opening, closing, headerless)
    reconciled = None
    if plan.residual is not None:
        reconciled = plan.residual <= 0.05

    owner_tax_id = _labeled_tax_id(preamble)
    kind = holder_kind(
        holder=holder,
        tax_id=owner_tax_id,
        header_cells=rows[header_index] if header_index is not None else (),
        requisites=preamble_text,
    )
    transactions, totals = _for_holder(plan.transactions, kind, owner_tax_id)

    metadata = StatementMetadata(
        source_filename=filename,
        title=title,
        parser_key="adaptive_bank_statement",
        account_holder=holder,
        account_number=account,
        currency=currency,
        period_start=period_start,
        period_end=period_end,
        opening_balance=opening,
        closing_balance=closing,
        transaction_count=len(transactions),
        totals=totals,
        reading_note=note,
        holder_kind=kind,
    )
    return ReadResult(
        statement=ParsedStatement(metadata=metadata, transactions=transactions),
        headerless=headerless,
        reconciled=reconciled,
    )


def _for_holder(
    transactions: list[StatementTransaction],
    kind: str | None,
    owner_tax_id: str | None,
) -> tuple[list[StatementTransaction], StatementTotals]:
    """У юрлица вид операции и итоги считаются так же, как у Kaspi Business.

    Слова физлица здесь врут: «Оплата по счёту № 26» от клиента ТОО — это
    поступление, а не «Покупка», и «Пополнений» у счёта с миллионом прихода
    выходило ноль.
    """
    if kind != LEGAL:
        return transactions, _totals(transactions)
    relabeled = [
        item.model_copy(update={"operation": legal_operation(item, owner_tax_id=owner_tax_id)})
        for item in transactions
    ]
    return relabeled, legal_totals(relabeled)


def _find_header(
    rows: list[list[object]],
) -> tuple[int | None, dict[str, int], list[int]]:
    best_index: int | None = None
    best_score = 0
    best_roles: dict[str, int] = {}
    best_extras: list[int] = []
    for index, row in enumerate(rows[:_HEADER_SCAN_ROWS]):
        roles, extras = _row_roles(row)
        score = _header_score(row, roles)
        if score > best_score:
            best_score = score
            best_index = index
            best_roles = roles
            best_extras = extras
    if best_index is None:
        return None, {}, []
    return best_index, best_roles, best_extras


def _row_roles(row: list[object]) -> tuple[dict[str, int], list[int]]:
    roles: dict[str, int] = {}
    extras: list[int] = []
    for index, cell in enumerate(row[:25]):
        role = _cell_role(cell)
        if role is None:
            continue
        if role == "date" and "date" in roles:
            role = "value_date"
        if role in {"detail", "counterparty"} and role in roles:
            extras.append(index)
            continue
        if role in roles:
            continue
        roles[role] = index
    return roles, extras


def _header_score(row: list[object], roles: dict[str, int]) -> int:
    if "date" not in roles:
        return 0
    if not any(role in roles for role in ("income", "expense", "amount")):
        return 0
    nonempty = sum(1 for cell in row[:25] if _normalize(cell))
    if nonempty < 3:
        return 0
    return len(roles) + nonempty


def _cell_role(value: object) -> str | None:
    text = _normalize(value).lower().replace("ё", "е")
    if not text or len(text) > 60:
        return None
    if any(marker in text for marker in _NOT_A_HEADER):
        return None
    if _DATE_FULL_RE.match(text) or _DATE_ISO_RE.match(text):
        return None
    if _coerce_amount(value) is not None and not any(char.isalpha() for char in text):
        return None
    for role, phrase in _PHRASES:
        if _contains_phrase(text, phrase):
            return role
    return None


def _contains_phrase(text: str, phrase: str) -> bool:
    if text == phrase:
        return True
    start = text.find(phrase)
    if start < 0:
        return False
    end = start + len(phrase)
    if start > 0 and text[start - 1].isalnum():
        return False
    if end < len(text) and text[end].isalnum():
        return False
    return True


def _layout_from_roles(header: list[object], roles: dict[str, int], extras: list[int]) -> Layout:
    money: list[MoneySlot] = []
    for role in ("income", "expense", "amount", "balance"):
        if role in roles:
            money.append(MoneySlot(roles[role], role, _normalize(header[roles[role]]) or role))
    detail_cols = [roles[role] for role in ("counterparty", "detail") if role in roles]
    detail_cols.extend(extras)
    return Layout(
        date_col=roles["date"],
        money=money,
        detail_cols=detail_cols,
        operation_col=roles.get("operation"),
        document_col=roles.get("document"),
        value_date_col=roles.get("value_date"),
        counterparty_col=roles.get("counterparty"),
        purpose_col=roles.get("detail"),
        headerless=False,
    )


def _infer_layout(rows: list[list[object]]) -> Layout | None:
    width = min(15, max((len(row) for row in rows[:80]), default=0))
    if width == 0:
        return None
    profiles: list[dict[str, float]] = []
    samples: list[list[float]] = []
    for index in range(width):
        dates = amounts = texts = seen = 0
        numbers: list[float] = []
        for row in rows[:80]:
            if index >= len(row) or row[index] in (None, ""):
                continue
            seen += 1
            if _coerce_date(row[index]) is not None:
                dates += 1
                continue
            amount = _coerce_amount(row[index])
            if amount is not None:
                amounts += 1
                numbers.append(amount)
                continue
            if len(_normalize(row[index])) >= 3:
                texts += 1
        profiles.append(
            {
                "dates": dates / seen if seen else 0.0,
                "amounts": amounts / seen if seen else 0.0,
                "texts": texts,
                "date_hits": dates,
                "amount_hits": amounts,
            }
        )
        samples.append(numbers)

    date_candidates = [
        index
        for index, profile in enumerate(profiles)
        if profile["date_hits"] >= 2 and profile["dates"] >= 0.45
    ]
    if not date_candidates:
        return None
    date_col = max(date_candidates, key=lambda index: profiles[index]["date_hits"])

    money_indexes = [
        index
        for index, profile in enumerate(profiles)
        if index != date_col and profile["amount_hits"] >= 2 and profile["amounts"] >= 0.4
    ]
    money_indexes = _drop_code_columns(money_indexes, samples)
    money_indexes = [index for index in money_indexes if not _is_serial_index(samples[index])]
    if not money_indexes:
        return None

    text_indexes = [
        index
        for index, profile in enumerate(profiles)
        if index != date_col and index not in money_indexes and profile["texts"] > 0
    ]
    text_indexes.sort(key=lambda index: profiles[index]["texts"], reverse=True)
    return Layout(
        date_col=date_col,
        money=[MoneySlot(index, "unknown", "") for index in money_indexes[:4]],
        detail_cols=text_indexes[:2],
        headerless=True,
    )


def _drop_code_columns(indexes: list[int], samples: list[list[float]]) -> list[int]:
    def is_code(values: list[float]) -> bool:
        if len(values) < 2:
            return False
        if any(abs(value) > 999 for value in values):
            return False
        return all(float(value).is_integer() for value in values)

    if not any(not is_code(samples[index]) for index in indexes):
        return indexes
    return [index for index in indexes if not is_code(samples[index])]


def _is_serial_index(values: list[float]) -> bool:
    integers = [int(value) for value in values if float(value).is_integer() and 0 < value < 100000]
    if len(integers) < 3 or len(integers) != len(values):
        return False
    start = integers[0]
    return integers == list(range(start, start + len(integers)))


def _merge_continuations(rows: list[list[object]], layout: Layout) -> list[list[object]]:
    money_indexes = [slot.index for slot in layout.money]
    merged: list[list[object]] = []
    # Строка закрыта, если за ней пошёл подвал или повтор шапки: иначе
    # «Обороты: Дебет Кредит За период…» с последней страницы дописывались
    # к последней операции.
    closed = False
    for row in rows:
        if not any(cell not in (None, "") for cell in row):
            continue
        if _cell_role(_cell(row, layout.date_col)) == "date":
            continue
        if _coerce_date(_cell(row, layout.date_col), allow_serial=not layout.headerless):
            merged.append(list(row))
            closed = False
            continue
        if not merged or closed:
            continue
        if _closes_table(row, layout):
            closed = True
            continue
        if any(_meaningful_amount(_cell(row, index)) for index in money_indexes):
            continue
        if not layout.detail_cols:
            continue
        current = merged[-1]
        if layout.headerless:
            extra = " ".join(
                _normalize(cell)
                for cell in row
                if _normalize(cell) and _coerce_amount(cell) is None and _coerce_date(cell) is None
            )
            if extra:
                _append_text(current, layout.detail_cols[0], extra)
            continue
        # Под шапкой у строки-продолжения каждая ячейка стоит под своей
        # колонкой, туда её и дописываем. Раньше всё склеивалось в первую
        # текстовую колонку, и в PDF Halyk контрагент перемешивался с
        # назначением построчно: «Товарищество с ограниченной по аренде, счет
        # на оплату № ответственностью "Алатау…». Числа в текстовой колонке —
        # часть текста («№ 34», «481981.86(KZT)»), их тоже не выбрасываем.
        text_cols = sorted(layout.detail_cols)
        for index, cell in enumerate(row):
            text = _normalize(cell)
            if not text:
                continue
            if index in layout.detail_cols:
                target = index
            else:
                if _coerce_amount(cell) is not None or _coerce_date(cell) is not None:
                    continue
                left = [column for column in text_cols if column < index]
                target = left[-1] if left else layout.detail_cols[0]
            _append_text(current, target, text)
    return merged


def _closes_table(row: list[object], layout: Layout) -> bool:
    """Подвал или повтор шапки, а не продолжение операции.

    «Обороты», «Итого» ищутся только под датой: в назначении платежа такие
    слова бывают. Слова в денежной колонке бывают только в шапке и в подвале
    («Дебет», «Кредит» над оборотами) — но это верно, лишь когда колонки
    известны по шапке.
    """
    under_date = _normalize(_cell(row, layout.date_col)).lower()
    if under_date.startswith(_TABLE_END_WORDS):
        return True
    if layout.headerless:
        return False
    for index in (slot.index for slot in layout.money):
        text = _normalize(_cell(row, index))
        if text and _coerce_amount(text) is None and any(char.isalpha() for char in text):
            return True
    return False


def _append_text(row: list[object], index: int, text: str) -> None:
    while len(row) <= index:
        row.append(None)
    previous = _normalize(row[index])
    row[index] = f"{previous} {text}".strip() if previous else text


def _raw_row(row: list[object], layout: Layout) -> RawRow | None:
    parsed_date = _coerce_date(_cell(row, layout.date_col), allow_serial=not layout.headerless)
    if parsed_date is None:
        return None
    detail_parts = [_normalize(_cell(row, index)) for index in layout.detail_cols]
    detail = " — ".join(part for part in detail_parts if part)
    operation_text = _normalize(_cell(row, layout.operation_col)) if layout.operation_col is not None else ""
    document = _normalize(_cell(row, layout.document_col)) if layout.document_col is not None else ""
    processing = None
    if layout.value_date_col is not None:
        processing = _coerce_date(_cell(row, layout.value_date_col), allow_serial=True)
    counterparty = _normalize(_cell(row, layout.counterparty_col)) if layout.counterparty_col is not None else ""
    purpose = _normalize(_cell(row, layout.purpose_col)) if layout.purpose_col is not None else ""
    return RawRow(
        date=parsed_date,
        values={slot.index: _coerce_amount(_cell(row, slot.index)) for slot in layout.money},
        detail=detail,
        operation_text=operation_text,
        document_number=document or None,
        processing_date=processing,
        counterparty=counterparty,
        purpose=purpose,
    )


def _choose_plan(
    raw_rows: list[RawRow],
    layout: Layout,
    opening: float | None,
    closing: float | None,
    headerless: bool,
) -> Plan | None:
    slots = layout.money
    by_role = {slot.role: slot for slot in slots}
    unknowns = [slot for slot in slots if slot.role == "unknown"]
    candidates: list[tuple[str, int, int | None, int | None, int | None, int | None, bool]] = []
    # mode, priority, income, expense, signed, balance, sign_from_balance

    income = by_role.get("income")
    expense = by_role.get("expense")
    amount = by_role.get("amount")
    balance = by_role.get("balance")

    if income and expense:
        candidates.append(("debit_credit", 0, income.index, expense.index, None, balance.index if balance else None, False))
        candidates.append(("debit_credit_swapped", 5, expense.index, income.index, None, balance.index if balance else None, False))
    elif income or expense:
        candidates.append((
            "headers",
            1,
            income.index if income else None,
            expense.index if expense else None,
            None,
            balance.index if balance else None,
            False,
        ))
    # Дебет и кредит уже говорят, что складывать. Отдельная «сумма» рядом с ними
    # не должна перебивать эту пару.
    if amount is not None and not (income and expense):
        candidates.append(("signed", 1, None, None, amount.index, balance.index if balance else None, False))
        if balance is not None:
            candidates.append(("balance_sign", 3, None, None, amount.index, balance.index, True))
    if len(unknowns) >= 2 and income is None and expense is None and amount is None:
        left, right = unknowns[0], unknowns[1]
        candidates.append(("guess_sides", 4, right.index, left.index, None, None, False))
        candidates.append(("guess_sides_swapped", 6, left.index, right.index, None, None, False))
        candidates.append(("balance_sign", 3, None, None, left.index, right.index, True))
        candidates.append(("balance_sign", 3, None, None, right.index, left.index, True))
    elif len(unknowns) == 1 and amount is None and income is None and expense is None:
        only = unknowns[0]
        if balance is not None:
            candidates.append(("signed", 2, None, None, only.index, balance.index, False))
            candidates.append(("balance_sign", 3, None, None, only.index, balance.index, True))
        else:
            candidates.append(("signed", 2, None, None, only.index, None, False))
    if balance is not None and amount is None and income is None and expense is None and not unknowns:
        candidates.append(("balance_delta", 4, None, None, None, balance.index, True))

    plans: list[Plan] = []
    for mode, priority, income_idx, expense_idx, signed_idx, balance_idx, sign_from_balance in candidates:
        transactions, consistency = _materialize(
            raw_rows,
            income_idx=income_idx,
            expense_idx=expense_idx,
            signed_idx=signed_idx,
            balance_idx=balance_idx,
            sign_from_balance=sign_from_balance,
            opening=opening,
        )
        if len(transactions) < 1:
            continue
        residual = _residual(opening, closing, transactions)
        plans.append(
            Plan(
                mode=mode,
                priority=priority,
                transactions=transactions,
                residual=residual,
                pseudo=_pseudo(mode, transactions, consistency),
                ignores_balance=balance_idx is not None and not sign_from_balance,
                all_positive=bool(transactions) and all(item.direction == "inflow" for item in transactions),
                income_label=_slot_label(slots, income_idx),
                expense_label=_slot_label(slots, expense_idx),
                amount_label=_slot_label(slots, signed_idx),
            )
        )
    if not plans:
        return None
    if any(plan.residual is not None for plan in plans):
        plans.sort(key=lambda plan: (plan.residual if plan.residual is not None else 1e18, plan.priority, -len(plan.transactions)))
    else:
        plans.sort(key=lambda plan: (plan.pseudo, plan.priority, -len(plan.transactions)))
    chosen = plans[0]
    chosen.transactions = _with_confidence(chosen.transactions, headerless, chosen.residual)
    return chosen


def _materialize(
    raw_rows: list[RawRow],
    *,
    income_idx: int | None,
    expense_idx: int | None,
    signed_idx: int | None,
    balance_idx: int | None,
    sign_from_balance: bool,
    opening: float | None,
) -> tuple[list[StatementTransaction], float]:
    transactions: list[StatementTransaction] = []
    previous = opening
    compared = 0
    matched = 0
    for raw in raw_rows:
        income: float | None = None
        expense: float | None = None
        if sign_from_balance and balance_idx is not None:
            current = raw.values.get(balance_idx)
            magnitude = raw.values.get(signed_idx) if signed_idx is not None else None
            if current is None:
                continue
            if previous is None:
                previous = current
                continue
            delta = round(current - previous, 2)
            if magnitude is not None and abs(magnitude) >= 0.005:
                compared += 1
                if abs(abs(magnitude) - abs(delta)) <= 1.0:
                    matched += 1
            previous = current
            if abs(delta) < 0.005:
                continue
            amount_abs = abs(delta)
            if magnitude is not None and abs(magnitude) >= 0.005 and abs(abs(magnitude) - abs(delta)) <= 1.0:
                amount_abs = abs(magnitude)
            if delta > 0:
                income = round(amount_abs, 2)
            else:
                expense = round(amount_abs, 2)
        elif signed_idx is not None:
            amount = raw.values.get(signed_idx)
            if amount is None or abs(amount) < 0.005:
                continue
            if amount > 0:
                income = round(amount, 2)
            else:
                expense = round(abs(amount), 2)
        else:
            income, expense = _sides(raw.values.get(income_idx), raw.values.get(expense_idx))
            if income is None and expense is None:
                continue
        transactions.append(_make_transaction(raw, income, expense))
    consistency = matched / compared if compared else 0.0
    return transactions, consistency


def _sides(income_value: float | None, expense_value: float | None) -> tuple[float | None, float | None]:
    income = income_value if income_value is not None and abs(income_value) >= 0.005 else None
    expense = expense_value if expense_value is not None and abs(expense_value) >= 0.005 else None
    if income is not None and income < 0 and expense is None:
        return None, round(abs(income), 2)
    if expense is not None and expense < 0 and income is None:
        return round(abs(expense), 2), None
    if income is not None:
        income = round(abs(income), 2)
    if expense is not None:
        expense = round(abs(expense), 2)
    return income, expense


def _make_transaction(raw: RawRow, income: float | None, expense: float | None) -> StatementTransaction:
    net = round((income or 0.0) - (expense or 0.0), 2)
    description = raw.detail.strip()
    operation = _classify(f"{description} {raw.operation_text}")
    if operation == "Операция":
        raw_operation = raw.operation_text.strip()
        if raw_operation and raw_operation.lower() not in _GENERIC_OPERATION_WORDS:
            operation = raw_operation
    if not description:
        description = operation
    return StatementTransaction(
        date=raw.date,
        amount=net,
        income=income,
        expense=expense,
        operation=operation,
        detail=description,
        details_operation=description,
        direction="inflow" if net > 0 else "outflow",
        document_number=raw.document_number,
        processing_date=raw.processing_date,
        # Контрагент, названный банком в своей колонке. «Финансы» берут его
        # в справочник контрагентов, вид «Юр счёт» — в колонку «Контрагент».
        raw_counterparty=raw.counterparty or None,
        comment=raw.purpose or None,
        source="adaptive",
    )


def _with_confidence(
    transactions: list[StatementTransaction],
    headerless: bool,
    residual: float | None,
) -> list[StatementTransaction]:
    if residual is not None and residual <= 0.05:
        confidence = 0.9
    elif headerless:
        confidence = 0.58
    else:
        confidence = 0.74
    return [item.model_copy(update={"source_confidence": confidence}) for item in transactions]


def _pseudo(mode: str, transactions: list[StatementTransaction], consistency: float) -> float:
    if mode == "balance_sign":
        return 0.05 if consistency >= 0.75 else 0.8
    if mode == "balance_delta":
        return 0.3 if len(transactions) >= 2 else 0.75
    if mode == "signed":
        return 0.15 if any(item.direction == "outflow" for item in transactions) else 0.55
    if mode in {"debit_credit", "headers"}:
        return 0.2 if _one_sided(transactions) >= 0.6 else 0.75
    if mode in {"debit_credit_swapped", "guess_sides_swapped"}:
        return 0.4 if _one_sided(transactions) >= 0.6 else 0.8
    if mode == "guess_sides":
        return 0.3 if _one_sided(transactions) >= 0.6 else 0.75
    return 0.5


def _one_sided(transactions: list[StatementTransaction]) -> float:
    if not transactions:
        return 0.0
    good = 0
    for item in transactions:
        has_income = item.income is not None
        has_expense = item.expense is not None
        if has_income != has_expense:
            good += 1
    return good / len(transactions)


def _residual(opening: float | None, closing: float | None, transactions: list[StatementTransaction]) -> float | None:
    if opening is None or closing is None:
        return None
    net = round(sum(item.amount for item in transactions), 2)
    return abs(round(opening + net - closing, 2))


def _slot_label(slots: list[MoneySlot], index: int | None) -> str:
    if index is None:
        return ""
    for slot in slots:
        if slot.index == index and slot.label:
            return slot.label
    return ""


def _compose_note(plan: Plan, opening: float | None, closing: float | None, headerless: bool) -> str:
    parts: list[str] = []
    if plan.mode == "debit_credit":
        parts.append("Дебет посчитан как расход, кредит — как приход.")
    elif plan.mode == "debit_credit_swapped":
        parts.append(
            "Колонки дебета и кредита поменяны местами: только так сходятся остаток на начало и на конец."
        )
    elif plan.mode == "headers":
        bits = []
        if plan.income_label:
            bits.append(f"в приход взята колонка «{plan.income_label}»")
        if plan.expense_label:
            bits.append(f"в расход — «{plan.expense_label}»")
        parts.append((", ".join(bits).capitalize() + ".") if bits else "Приход и расход взяты из подписанных колонок.")
    elif plan.mode == "signed":
        label = plan.amount_label or "Сумма"
        parts.append(f"Колонка «{label}» со знаком: плюс — приход, минус — расход.")
    elif plan.mode == "balance_sign":
        parts.append(
            "Суммы без знака. Направление взято по изменению остатка, сама колонка остатка в обороты не входит."
        )
    elif plan.mode == "balance_delta":
        parts.append("Сумма операции посчитана как изменение остатка от строки к строке.")
    elif plan.mode == "guess_sides":
        parts.append("Левая денежная колонка посчитана как расход, правая — как приход.")
    elif plan.mode == "guess_sides_swapped":
        parts.append("Правая денежная колонка посчитана как расход, левая — как приход: так сходится остаток.")
    if plan.ignores_balance and plan.mode not in {"balance_sign", "balance_delta"}:
        parts.append("Колонка остатка в обороты не входит.")
    if headerless:
        parts.append("Подписи колонок были неочевидны, поэтому колонки выбраны по датам и числам в ячейках.")

    income_total = round(sum(item.income or 0.0 for item in plan.transactions), 2)
    expense_total = round(sum(item.expense or 0.0 for item in plan.transactions), 2)
    count = len(plan.transactions)
    parts.append(
        f"{count} {_operations_word(count)}: приход {_money(income_total)}, расход {_money(expense_total)}."
    )
    if plan.residual is not None and opening is not None and closing is not None:
        if plan.residual <= 0.05:
            parts.append("Остаток на конец сходится с суммой операций.")
        else:
            parts.append(
                f"Остаток на конец не сходится на {_money(plan.residual)}. "
                "Проверьте, все ли строки попали в расчёт."
            )
    elif plan.mode == "signed" and plan.all_positive:
        parts.append("Все суммы пришли без минуса и без колонки расхода, поэтому они записаны как приход.")
    return " ".join(parts)


def _classify(text: str) -> str:
    folded = text.lower().replace("ё", "е")
    groups = (
        ("Снятие", ("снятие", "банкомат", "atm", "cash withdrawal")),
        ("Пополнение", ("пополнение", "зачисление", "зарплат")),
        ("Перевод", ("перевод", "transfer")),
        ("Покупка", ("покуп", "оплата", "pos ", "магазин")),
    )
    for label, markers in groups:
        if any(marker in folded for marker in markers):
            return label
    return "Операция"


def _totals(transactions: list[StatementTransaction]) -> StatementTotals:
    buckets: dict[str, float] = defaultdict(float)
    for item in transactions:
        if item.income is not None:
            buckets["income_total"] += item.income
        if item.expense is not None:
            buckets["expense_total"] += item.expense
        if item.operation == "Покупка" and item.expense is not None:
            buckets["purchase_total"] += item.expense
        if item.operation == "Перевод" and item.expense is not None:
            buckets["transfer_total"] += item.expense
        if item.operation == "Пополнение" and item.income is not None:
            buckets["topup_total"] += item.income
        if item.operation == "Снятие" and item.expense is not None:
            buckets["cash_withdrawal_total"] += item.expense
    return StatementTotals(**{key: round(value, 2) for key, value in buckets.items()})


def _operations_word(count: int) -> str:
    mod100 = count % 100
    if 11 <= mod100 <= 14:
        return "операций"
    mod10 = count % 10
    if mod10 == 1:
        return "операция"
    if 2 <= mod10 <= 4:
        return "операции"
    return "операций"


def _money(value: float) -> str:
    sign = "-" if value < 0 else ""
    whole, frac = f"{abs(value):.2f}".split(".")
    groups: list[str] = []
    while whole:
        groups.append(whole[-3:])
        whole = whole[:-3]
    return sign + " ".join(reversed(groups)) + "," + frac


def _labeled_amount(rows: list[list[object]], labels: tuple[str, ...]) -> float | None:
    for row in rows:
        for index, cell in enumerate(row):
            text = _normalize(cell).lower().replace("ё", "е").rstrip(":")
            if not text or len(text) > 80:
                continue
            if not any(text == label or text.startswith(label) for label in labels):
                continue
            tail = _normalize(cell)
            if ":" in tail:
                amount = _coerce_amount(tail.split(":", 1)[1])
                if amount is not None:
                    return amount
            for follower in row[index + 1 : index + 4]:
                amount = _coerce_amount(follower)
                if amount is not None:
                    return amount
    return None


def _labeled_text(rows: list[list[object]], labels: tuple[str, ...]) -> str | None:
    for row in rows:
        for index, cell in enumerate(row):
            text = _normalize(cell).lower().replace("ё", "е").rstrip(":")
            if not text or len(text) > 80:
                continue
            if not any(text == label or text.startswith(label) for label in labels):
                continue
            raw = _normalize(cell)
            if ":" in raw:
                tail = _normalize(raw.split(":", 1)[1])
                if tail and _coerce_amount(tail) is None:
                    return tail[:120]
            for follower in row[index + 1 : index + 4]:
                value = _normalize(follower)
                if value and _coerce_amount(follower) is None and _coerce_date(follower) is None:
                    return value[:120]
    return None


def _labeled_tax_id(rows: list[list[object]]) -> str | None:
    """БИН или ИИН владельца из реквизитов: «ИИН/БИН | 211240002990».

    Отдельно от `_labeled_text`: двенадцать цифр та читает как сумму и
    пропускает.
    """
    for row in rows:
        for index, cell in enumerate(row):
            text = _normalize(cell)
            label = _TAX_ID_LABEL_RE.match(text)
            if not label:
                continue
            for candidate in [label.group(1), *row[index + 1 : index + 4]]:
                found = _TAX_ID_VALUE_RE.search(re.sub(r"[\s\xa0]", "", _normalize(candidate)))
                if found:
                    return found.group(1)
    return None


def _amount_after_label(text: str, labels: tuple[str, ...]) -> float | None:
    # Часть PDF отдаёт пробелы неразрывными: «Исходящий\xa0остаток:» иначе
    # не узнаётся, и сверка с банком молча пропадает.
    text = text.replace("\xa0", " ")
    folded = text.lower().replace("ё", "е")
    for label in labels:
        start = 0
        while True:
            found = folded.find(label, start)
            if found < 0:
                break
            window = text[found + len(label) : found + len(label) + 48]
            amount = _first_amount_token(window)
            if amount is not None:
                return amount
            start = found + len(label)
    return None


def _text_after_label(text: str, labels: tuple[str, ...]) -> str | None:
    for line in text.splitlines():
        found = _labeled_text([[line]], labels)
        if found:
            return found
    return None


def _period(text: str) -> tuple[str | None, str | None]:
    match = _PERIOD_FROM_TO_RE.search(text) or _PERIOD_DASH_RE.search(text)
    if not match:
        return None, None
    start = _coerce_date(match.group(1))
    end = _coerce_date(match.group(2))
    return start, end


def _currency(text: str) -> str | None:
    explicit = re.search(r"валюта(?:\s+счета)?\s*[:\-]?\s*([A-Za-z]{3}|тенге)", text, re.IGNORECASE)
    if explicit:
        token = explicit.group(1).lower()
        return "KZT" if token == "тенге" else token.upper()
    if re.search(r"\bKZT\b|₸|тенге", text, re.IGNORECASE):
        return "KZT"
    for code in ("USD", "EUR", "RUB", "GBP", "CNY"):
        if re.search(rf"\b{code}\b", text):
            return code
    return None


def _iban(text: str) -> str | None:
    match = _IBAN_RE.search(text.replace(" ", ""))
    if match:
        return match.group(0).upper()
    return None


def _bank_name(text: str) -> str | None:
    folded = text.lower().replace("ё", "е")
    for needle, name in _BANKS:
        if needle in folded:
            return name
    return None


def _first_amount_token(text: str) -> float | None:
    for match in _AMOUNT_TOKEN_RE.finditer(text):
        window = text[max(0, match.start() - 1) : match.end() + 6]
        if _DATE_IN_TEXT_RE.search(window):
            continue
        amount = _coerce_amount(match.group(0))
        if amount is not None:
            return amount
    return None


def _meaningful_amount(value: object) -> bool:
    amount = _coerce_amount(value)
    return amount is not None and abs(amount) >= 0.005


def _coerce_amount(value: object) -> float | None:
    if value is None or isinstance(value, bool) or isinstance(value, datetime):
        return None
    if isinstance(value, (int, float)):
        if isinstance(value, float) and value != value:
            return None
        return round(float(value), 2)
    text = _normalize(value).replace("−", "-")
    if not text or text in {"-", "—", "–"}:
        return None
    if _DATE_FULL_RE.match(text) or _DATE_ISO_RE.match(text):
        return None
    negative = text.startswith("-") or (text.startswith("(") and text.endswith(")"))
    cleaned = _CURRENCY_RE.sub("", text)
    cleaned = cleaned.replace("−", "-").strip().strip("()")
    cleaned = cleaned.replace("\u00a0", "").replace(" ", "")
    cleaned = cleaned.lstrip("+-")
    if not cleaned or any(char.isalpha() for char in cleaned):
        return None
    if cleaned.count(",") and cleaned.count("."):
        if cleaned.rfind(",") > cleaned.rfind("."):
            cleaned = cleaned.replace(".", "").replace(",", ".")
        else:
            cleaned = cleaned.replace(",", "")
    elif "," in cleaned:
        left, right = cleaned.split(",", 1)
        if right.isdigit() and len(right) <= 2:
            cleaned = f"{left}.{right}"
        elif right.isdigit() and len(right) == 3:
            cleaned = left + right
        else:
            return None
    elif cleaned.count(".") > 1:
        cleaned = cleaned.replace(".", "")
    elif cleaned.count(".") == 1:
        left, right = cleaned.split(".", 1)
        if right.isdigit() and len(right) == 3 and left.isdigit():
            cleaned = left + right
    if not re.fullmatch(r"\d+(?:\.\d+)?", cleaned):
        return None
    amount = round(float(cleaned), 2)
    return -amount if negative else amount


def _coerce_date(value: object, allow_serial: bool = False) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, date):
        return value.strftime("%d.%m.%Y")
    if allow_serial and isinstance(value, (int, float)) and not isinstance(value, bool):
        serial = float(value)
        if 30000 <= serial <= 80000:
            parsed = _EXCEL_EPOCH + timedelta(days=int(serial))
            return parsed.strftime("%d.%m.%Y")
        return None
    text = _normalize(value)
    matched = _DATE_FULL_RE.match(text)
    if matched:
        day, month, year = matched.groups()
        if len(year) == 2:
            year = ("20" if int(year) < 70 else "19") + year
        try:
            datetime(int(year), int(month), int(day))
        except ValueError:
            return None
        return f"{int(day):02d}.{int(month):02d}.{year}"
    iso = _DATE_ISO_RE.match(text)
    if iso:
        year, month, day = iso.groups()
        try:
            datetime(int(year), int(month), int(day))
        except ValueError:
            return None
        return f"{day}.{month}.{year}"
    return None


def _parse_stamp(value: str) -> datetime | None:
    try:
        return datetime.strptime(value, "%d.%m.%Y")
    except ValueError:
        return None


def _cell(row: list[object], index: int | None) -> object | None:
    if index is None or index >= len(row):
        return None
    return row[index]


def _normalize(value: object) -> str:
    if value is None:
        return ""
    text = str(value).replace("\xa0", " ").replace("\n", " ")
    # PyMuPDF с кириллическим шрифтом иногда отдаёт минус как мягкий перенос.
    for dash in ("\u00ad", "\u2212", "\u2012", "\u2013", "\u2014"):
        text = text.replace(dash, "-")
    text = unicodedata.normalize("NFC", text)
    return re.sub(r"[ \t]+", " ", text).strip()


def _rows_text(rows: list[list[object]]) -> str:
    lines = []
    for row in rows:
        parts = [_normalize(cell) for cell in row if _normalize(cell)]
        if parts:
            lines.append(" ".join(parts))
    return "\n".join(lines)


def _specialist_fingerprint(filename: str, content: bytes) -> bool:
    extension = Path(filename).suffix.lower()
    try:
        if extension == ".pdf":
            document = fitz.open(stream=content, filetype="pdf")
            try:
                text = "\n".join(document[index].get_text("text") for index in range(min(2, document.page_count)))
            finally:
                document.close()
            if "Kaspi Gold" in text and "ВЫПИСКА" in text:
                return True
            if looks_like_halyk_personal(text):
                return True
            return False
        if extension not in {".xlsx", ".xlsm"}:
            return False
        sheets = _excel_sheets(content)
    except Exception:
        return False
    for rows in sheets:
        sample = _rows_text(rows[:20])
        if "Kaspi Gold" in sample and "ВЫПИСКА" in sample:
            return True
        folded = sample.lower()
        if "kaspi" not in folded and "текущий счет" not in folded and "текущий счёт" not in folded:
            continue
        for row in rows[:20]:
            cells = [_normalize(cell).lower() for cell in row if _normalize(cell)]
            if (
                any("документ" in cell for cell in cells)
                and any("дата операции" in cell for cell in cells)
                and any(cell == "дебет" for cell in cells)
                and any(cell == "кредит" for cell in cells)
            ):
                return True
    return False


def _pdf_grid(content: bytes) -> tuple[list[list[object]], str] | None:
    document = fitz.open(stream=content, filetype="pdf")
    try:
        if document.page_count == 0:
            return None
        page_rows: list[list[tuple]] = []
        texts: list[str] = []
        rulings: list[list[tuple[float, float, float]]] = []
        for index, page in enumerate(document):
            if index >= _MAX_PDF_PAGES:
                break
            words = [word for word in page.get_text("words") if str(word[4]).strip()]
            page_rows.append(words)
            texts.append(page.get_text("text"))
            rulings.append(_vertical_rulings(page))
    finally:
        document.close()
    if not any(page_rows):
        return None

    visual_pages = [_visual_rows(words) for words in page_rows]
    header = _find_visual_header(visual_pages)
    if header is None:
        flat = [[cell for cell in row.cells] for page in visual_pages for row in page]
        return flat, "\n".join(texts)

    _, header_row, bounds = header
    header_page = next(
        index for index, page in enumerate(visual_pages) if any(row is header_row for row in page)
    )
    grid: list[list[object]] = []
    active_bounds = _snap_bounds(bounds, header_row, rulings[header_page])
    reached = False
    for page_index, page in enumerate(visual_pages):
        for row in page:
            if row is header_row:
                reached = True
                grid.append(list(row.cells))
                continue
            if not reached:
                # Реквизиты над таблицей режутся по своим промежуткам, а не по
                # колонкам таблицы: иначе «Клиент  ТОО Omar Development and
                # Consulting» разъезжался на «ТОО Omar | Development and |
                # Consulting», и владельцем счёта становился «ТОО Omar».
                grid.append(list(row.cells))
                continue
            if _cell_role(row.cells[0] if row.cells else "") == "date" and len(row.cells) >= 3:
                refreshed = _bounds_from_groups(row.groups, row.page_width)
                if refreshed:
                    active_bounds = _snap_bounds(refreshed, row, rulings[page_index])
                    grid.append(list(row.cells))
                    continue
            if active_bounds:
                grid.append(_row_by_bounds(row.words, active_bounds))
            else:
                grid.append(list(row.cells))
    return grid, "\n".join(texts)


def _vertical_rulings(page: fitz.Page) -> list[tuple[float, float, float]]:
    """Вертикальные линейки таблицы: (x, верх, низ)."""
    found: list[tuple[float, float, float]] = []
    try:
        drawings = page.get_drawings()
    except Exception:
        return found
    for drawing in drawings:
        for item in drawing.get("items", ()):
            if item[0] == "l":
                start, end = item[1], item[2]
                if abs(start.x - end.x) <= 1 and abs(start.y - end.y) >= 2:
                    found.append(((start.x + end.x) / 2, min(start.y, end.y), max(start.y, end.y)))
            elif item[0] == "re":
                rect = item[1]
                if rect.height < 2:
                    continue
                if rect.width <= 2:
                    found.append(((rect.x0 + rect.x1) / 2, rect.y0, rect.y1))
                else:
                    # Ячейка, нарисованная прямоугольником: её края — те же линейки.
                    found.append((rect.x0, rect.y0, rect.y1))
                    found.append((rect.x1, rect.y0, rect.y1))
    return found


def _snap_bounds(
    bounds: list[tuple[float, float]],
    header: VisualRow,
    rulings: list[tuple[float, float, float]],
) -> list[tuple[float, float]]:
    """Границы колонок по линейкам таблицы, если они нарисованы.

    Без линеек граница — середина между подписями шапки. Подписи стоят по
    центру колонок, и длинный текст широкой колонки заезжал в соседнюю: в
    выписке Halyk «счет на оплату № 34» терял «34» в колонке НДС.
    """
    if not bounds or not rulings or len(header.groups) != len(bounds):
        return bounds
    top = min(word[1] for word in header.words)
    bottom = max(word[3] for word in header.words)
    xs = sorted({round(x, 1) for x, y0, y1 in rulings if y0 <= bottom + 2 and y1 >= top - 2})
    if len(xs) < 2:
        return bounds
    cells: list[tuple[float, float]] = []
    for group in header.groups:
        left_edge = min(word[0] for word in group)
        right_edge = max(word[2] for word in group)
        lefts = [x for x in xs if x <= left_edge + 1]
        rights = [x for x in xs if x >= right_edge - 1]
        if not lefts or not rights:
            return bounds
        cells.append((lefts[-1], rights[0]))
    # Две подписи в одной ячейке — шапка разбита не по колонкам, линейкам не верим.
    if any(right > next_left + 1 for (_, right), (next_left, _) in zip(cells, cells[1:])):
        return bounds
    snapped: list[tuple[float, float]] = []
    for index, (left, _right) in enumerate(cells):
        low = 0.0 if index == 0 else left
        high = bounds[-1][1] if index == len(cells) - 1 else cells[index + 1][0]
        snapped.append((low, high))
    return snapped


@dataclass
class VisualRow:
    cells: list[str]
    groups: list[list[tuple]]
    words: list[tuple]
    page_width: float


def _visual_rows(words: list[tuple]) -> list[VisualRow]:
    if not words:
        return []
    heights = sorted(word[3] - word[1] for word in words)
    median_height = heights[len(heights) // 2] or 8
    tolerance = max(2.0, median_height * 0.55)
    ordered = sorted(words, key=lambda word: ((word[1] + word[3]) / 2, word[0]))
    clustered: list[tuple[float, list[tuple]]] = []
    for word in ordered:
        center = (word[1] + word[3]) / 2
        if clustered and abs(center - clustered[-1][0]) <= tolerance:
            clustered[-1][1].append(word)
        else:
            clustered.append((center, [word]))
    page_width = max(word[2] for word in words) + 8
    rows: list[VisualRow] = []
    for _, group in clustered:
        groups = _split_words(group)
        rows.append(
            VisualRow(
                cells=[" ".join(str(word[4]) for word in cell_words) for cell_words in groups],
                groups=groups,
                words=group,
                page_width=page_width,
            )
        )
    return rows


def _split_words(words: list[tuple]) -> list[list[tuple]]:
    ordered = sorted(words, key=lambda word: word[0])
    widths = [
        (word[2] - word[0]) / max(len(str(word[4])), 1)
        for word in ordered
        if str(word[4]).strip()
    ]
    widths.sort()
    char_width = widths[len(widths) // 2] if widths else 5
    threshold = max(10.0, char_width * 1.8)
    groups: list[list[tuple]] = [[ordered[0]]]
    for previous, word in zip(ordered, ordered[1:]):
        gap = word[0] - previous[2]
        if gap > threshold:
            groups.append([word])
        else:
            groups[-1].append(word)
    return groups


def _find_visual_header(
    pages: list[list[VisualRow]],
) -> tuple[int, VisualRow, list[tuple[float, float]]] | None:
    best: tuple[int, int, VisualRow, list[tuple[float, float]]] | None = None
    seen = 0
    for page in pages:
        for row in page:
            seen += 1
            if seen > _HEADER_SCAN_ROWS:
                break
            roles, _extras = _row_roles(row.cells)
            score = _header_score(row.cells, roles)
            if score <= 0:
                continue
            bounds = _bounds_from_groups(row.groups, row.page_width)
            if best is None or score > best[0]:
                best = (score, seen, row, bounds)
        if seen > _HEADER_SCAN_ROWS:
            break
    if best is None:
        return None
    return best[1], best[2], best[3]


def _bounds_from_groups(groups: list[list[tuple]], page_width: float) -> list[tuple[float, float]]:
    spans = []
    for group in groups:
        if not group:
            continue
        spans.append((min(word[0] for word in group), max(word[2] for word in group)))
    bounds: list[tuple[float, float]] = []
    for index, (left_edge, right_edge) in enumerate(spans):
        left = 0.0 if index == 0 else (spans[index - 1][1] + left_edge) / 2
        if index == len(spans) - 1:
            right = max(page_width, right_edge + 4)
        else:
            right = (right_edge + spans[index + 1][0]) / 2
        bounds.append((left, right))
    return bounds


def _row_by_bounds(words: list[tuple], bounds: list[tuple[float, float]]) -> list[str]:
    buckets: list[list[tuple]] = [[] for _ in bounds]
    for word in words:
        center = (word[0] + word[2]) / 2
        placed = False
        for index, (left, right) in enumerate(bounds):
            if left <= center < right:
                buckets[index].append(word)
                placed = True
                break
        if not placed and bounds and center >= bounds[-1][0]:
            buckets[-1].append(word)
    return [
        " ".join(str(word[4]) for word in sorted(bucket, key=lambda item: item[0])).strip()
        for bucket in buckets
    ]


def _read_pdf_lines(filename: str, content: bytes) -> ReadResult | None:
    document = fitz.open(stream=content, filetype="pdf")
    try:
        text = "\n".join(
            document[index].get_text("text")
            for index in range(min(document.page_count, _MAX_PDF_PAGES))
        )
    finally:
        document.close()
    lines = [_normalize(line) for line in text.splitlines()]
    lines = [line for line in lines if line]
    blocks: list[tuple[str, str]] = []
    current_date: str | None = None
    current_parts: list[str] = []

    def flush() -> None:
        if current_date is not None:
            blocks.append((current_date, " ".join(current_parts).strip()))

    for line in lines:
        matched = _LINE_DATE_RE.match(line)
        parsed = _coerce_date(matched.group(1)) if matched else None
        if parsed is not None and matched is not None:
            flush()
            current_date = parsed
            current_parts = [matched.group(2)] if matched.group(2) else []
            continue
        if current_date is not None and not _cell_role(line):
            current_parts.append(line)
    flush()

    raw_rows: list[RawRow] = []
    for parsed_date, block in blocks:
        amount = _first_amount_token(block)
        if amount is None or abs(amount) < 0.005:
            continue
        description = _AMOUNT_TOKEN_RE.sub(" ", block)
        description = _normalize(description).strip(" -—|")
        raw_rows.append(
            RawRow(
                date=parsed_date,
                values={0: amount},
                detail=description,
                operation_text="",
                document_number=None,
                processing_date=None,
            )
        )
    if len(raw_rows) < 1:
        return None
    opening = _amount_after_label(text, _OPENING_LABELS)
    closing = _amount_after_label(text, _CLOSING_LABELS)
    layout = Layout(date_col=0, money=[MoneySlot(0, "amount", "Сумма")], detail_cols=[1], headerless=True)
    plan = _choose_plan(raw_rows, layout, opening, closing, headerless=True)
    if plan is None:
        return None
    period_start, period_end = _period(text)
    bank = _bank_name(f"{filename}\n{text[:1500]}")
    # Уверенность остаётся осторожной: таблицы не было. В пояснении не пишем,
    # что путались в заголовках, — их просто не было, строки читались по дате.
    note = "Строки выписки прочитаны по дате в начале строки. " + _compose_note(
        plan, opening, closing, headerless=False
    )
    holder = _text_after_label(text, _HOLDER_LABELS)
    first_date = next((index for index, line in enumerate(lines) if _LINE_DATE_RE.match(line)), len(lines))
    kind = holder_kind(holder=holder, tax_id=None, requisites="\n".join(lines[:first_date]))
    transactions, totals = _for_holder(plan.transactions, kind, None)
    metadata = StatementMetadata(
        source_filename=filename,
        title=f"Выписка {bank}" if bank else "Банковская выписка",
        parser_key="adaptive_bank_statement",
        account_holder=holder,
        account_number=_iban(text),
        currency=_currency(text),
        period_start=period_start,
        period_end=period_end,
        opening_balance=opening,
        closing_balance=closing,
        transaction_count=len(transactions),
        totals=totals,
        reading_note=note,
        holder_kind=kind,
    )
    reconciled = plan.residual <= 0.05 if plan.residual is not None else None
    return ReadResult(
        statement=ParsedStatement(metadata=metadata, transactions=transactions),
        headerless=True,
        reconciled=reconciled,
    )
