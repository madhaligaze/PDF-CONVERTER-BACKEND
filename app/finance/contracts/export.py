"""Выгрузка реестра в .xlsx — теми же листами, блоками и шапками, что в файле.

Страховка на переходный период и для тех, кому нужен файл. Лист выгрузки —
это отбор реестра: договор стоит в первом подходящем блоке листа, как стоял бы
в Excel. Шапка блока — подписи колонок из файла (у «Заказчик ГК / Заказчик ГК»
заказчик снова в F), текст соглашений — как был, даты — датами Excel, суммы —
числами. Цвета шапки исходного файла возвращаются в выгрузке: на экране их нет,
а в файле человек их ждёт.
"""
from __future__ import annotations

import io
import uuid
from datetime import date
from decimal import Decimal
from typing import Any, Sequence

import sqlalchemy as sa
from sqlalchemy.orm import Session

from app.finance import history
from app.finance.contracts import views as views_module
from app.finance.contracts.models import Contract, EntityView
from app.finance.contracts.service import (
    Access,
    Actor,
    Registry,
    facts_of,
    people_of,
    visible_to,
)
from app.finance.models import Workspace

ROW_NUMBER = "row_number"
#: Колонки главного листа, если у листа своих колонок нет (реестр заведён
#: не из файла).
DEFAULT_KEYS = (
    ROW_NUMBER, "status", "folder_url", "planned_end_at", "people", "executor", "customer",
    "number", "signed_at", "department", "type", "subject", "amount", "paid_snapshot",
    "remaining_snapshot", "amendments_text", "amendments_summary_text", "end_date", "note",
)


def _text(registry: Registry, contract: Contract, key: str, people: Sequence[uuid.UUID]) -> Any:
    """Значение ячейки выгрузки: так, как его пишут в реестре."""
    if key in ("executor", "customer"):
        party_id = contract.executor_id if key == "executor" else contract.customer_id
        party = registry.parties.get(party_id) if party_id else None
        return party.name if party else None
    if key in ("type", "subject", "status", "economic_role"):
        value_id = getattr(contract, f"{key}_id")
        value = registry.values.get(value_id) if value_id else None
        return value.value if value else None
    if key == "department":
        department = registry.departments.get(contract.department_id) if contract.department_id else None
        return department.code if department else None
    if key == "people":
        names = [registry.employees[item].full_name for item in people if item in registry.employees]
        return ", ".join(names) or None
    if key == "amount":
        if contract.amount is not None:
            return float(contract.amount)
        return contract.amount_terms or None
    if key in ("paid_snapshot", "remaining_snapshot"):
        raw = (contract.file_snapshot or {}).get(key.replace("_snapshot", ""))
        try:
            return float(Decimal(str(raw))) if raw not in (None, "") else None
        except ArithmeticError:
            return raw
    if key in ("signed_at", "planned_end_at", "end_date"):
        return getattr(contract, key)
    if key in ("billing", "end_kind"):
        return getattr(contract, key) or None
    if key in ("number", "folder_url", "amendments_text", "amendments_summary_text", "note", "amount_terms", "currency"):
        return getattr(contract, key) or None
    raw = (contract.attrs or {}).get(key)
    field = registry.field_by_key.get(key)
    if raw in (None, "", []) or field is None:
        return None
    try:
        if field.type in ("list",):
            value = registry.values.get(uuid.UUID(str(raw)))
            return value.value if value else raw
        if field.type == "multi_list":
            return ", ".join(
                registry.values[uuid.UUID(item)].value for item in raw if uuid.UUID(item) in registry.values
            )
        if field.type == "person":
            return ", ".join(
                registry.employees[uuid.UUID(item)].full_name for item in raw if uuid.UUID(item) in registry.employees
            )
        if field.type in ("number", "money"):
            return float(Decimal(str(raw)))
        if field.type == "date":
            return date.fromisoformat(str(raw))
    except (ValueError, KeyError, ArithmeticError):
        return str(raw)
    return raw


def build(session: Session, workspace: Workspace, access: Access, actor: Actor, view_keys: Sequence[str] | None = None) -> bytes:
    """Книга .xlsx: лист реестра — лист книги, блок — название, шапка и строки.

    Книга пишется потоком (`write_only`): строки уходят в файл по мере записи,
    а не копятся объектами ячеек. Обычная книга держала в памяти каждую ячейку
    со стилем — выгрузка 20 000 договоров поднимала процесс на 380 МБ, и
    память обратно не возвращалась. В потоковой книге ширины колонок и
    закрепление шапки задаются до первой строки: их знают заранее из листа.
    """
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    registry = Registry(session, workspace)
    contracts = list(
        session.scalars(
            sa.select(Contract)
            .where(Contract.workspace_id == workspace.id, Contract.deleted_at.is_(None))
            .order_by(Contract.position, Contract.created_at)
        )
    )
    people = people_of(session, [item.id for item in contracts])
    visible = [item for item in contracts if visible_to(item, registry, access, people.get(item.id, []))]
    # Стороны всех выгружаемых договоров — одним запросом, а не по одному на
    # договор внутри facts_of.
    registry.parties_for({pid for item in visible for pid in (item.executor_id, item.customer_id)})
    facts = {item.id: facts_of(item, registry, people.get(item.id, [])) for item in visible}
    hidden = set(access.hidden)

    views: list[EntityView] = [
        view for view in registry.views if not view_keys or view.key in view_keys
    ]
    book = Workbook(write_only=True)
    bold = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")
    taken_titles: list[str] = []
    for view in views:
        title = _sheet_title(view.title, taken_titles)
        taken_titles.append(title)
        sheet = book.create_sheet(title=title)
        fill_hex = (view.style or {}).get("header_fill") or ""
        fill = PatternFill("solid", start_color=fill_hex) if fill_hex else None
        blocks = list(view.blocks or [])
        layouts = [
            [column for column in (block.get("columns") or []) if column.get("key") not in hidden]
            or _default_columns(registry, hidden)
            for block in blocks
        ]
        # Договор — в одном блоке листа: в первом подходящем (`views.place`).
        members: list[list[Contract]] = [[] for _ in blocks]
        for item in visible:
            index = views_module.place(view, facts[item.id])
            if index is not None and index < len(members):
                members[index].append(item)
        widths: dict[int, float] = {}
        for columns in layouts:
            for c_index, column in enumerate(columns, start=1):
                if column.get("width"):
                    widths[c_index] = max(widths.get(c_index, 0), float(column["width"]))
        for c_index, width in widths.items():
            sheet.column_dimensions[get_column_letter(c_index)].width = width
        if len(blocks) == 1:
            header_row = 2 if blocks[0].get("title") else 1
            sheet.freeze_panes = f"A{header_row + 1}"

        def styled(value: Any, *, font=None, alignment=None, number_format=None, cell_fill=None) -> WriteOnlyCell:
            cell = WriteOnlyCell(sheet, value=value)
            if font is not None:
                cell.font = font
            if alignment is not None:
                cell.alignment = alignment
            if number_format is not None:
                cell.number_format = number_format
            if cell_fill is not None:
                cell.fill = cell_fill
            return cell

        for number, block in enumerate(blocks):
            columns = layouts[number]
            if number:
                sheet.append([])
                sheet.append([])
            if block.get("title"):
                sheet.append([None, styled(block["title"], font=bold)])
            sheet.append([
                styled(column.get("label") or _title(registry, column["key"]), font=bold, alignment=wrap, cell_fill=fill)
                for column in columns
            ])
            for position, item in enumerate(members[number], start=1):
                mine = people.get(item.id, [])
                row: list[Any] = []
                for column in columns:
                    key = column["key"]
                    if key == ROW_NUMBER:
                        row.append(position)
                        continue
                    value = _text(registry, item, key, mine)
                    if isinstance(value, date):
                        row.append(styled(value, number_format="DD.MM.YYYY"))
                    elif isinstance(value, float):
                        row.append(styled(value, number_format="# ##0"))
                    elif isinstance(value, str) and "\n" in value:
                        row.append(styled(value, alignment=wrap))
                    else:
                        row.append(value)
                sheet.append(row)
    if not views:
        book.create_sheet("Реестр")
    buffer = io.BytesIO()
    book.save(buffer)
    history.write(
        session,
        workspace,
        kind="contract.export",
        entity="contract",
        title=f"реестр выгружен в .xlsx: {len(visible)} договоров, листов {len(views)}",
        after={"views": [view.key for view in views], "contracts": len(visible)},
        actor=actor.email,
    )
    return buffer.getvalue()


def _default_columns(registry: Registry, hidden: set[str]) -> list[dict[str, Any]]:
    return [
        {"key": key, "label": "№" if key == ROW_NUMBER else _title(registry, key)}
        for key in DEFAULT_KEYS
        if key == ROW_NUMBER or (key in registry.field_by_key and key not in hidden)
    ]


def _title(registry: Registry, key: str) -> str:
    if key == ROW_NUMBER:
        return "№"
    field = registry.field_by_key.get(key)
    return field.title if field else key


def _sheet_title(title: str, taken: list[str]) -> str:
    clean = "".join(ch for ch in (title or "Лист") if ch not in "[]:*?/\\")[:31] or "Лист"
    candidate, index = clean, 2
    while candidate in taken:
        suffix = f" {index}"
        candidate = clean[: 31 - len(suffix)] + suffix
        index += 1
    return candidate


__all__ = ["build"]
