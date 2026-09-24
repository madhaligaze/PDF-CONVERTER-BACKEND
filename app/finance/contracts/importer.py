"""Загрузка реестра из Excel: протокол разбора из девяти пунктов.

Что делает разбор
─────────────────
1. **Блоки.** Лист режется на блоки: строка названия («АРЕНДА») → строка
   шапки → строки договоров. У каждого блока своя шапка: в «Заказчик ГК»
   соседние блоки держат стороны в колонках наоборот, и одна шапка на лист
   перевернула бы стороны.
2. **Колонки** сопоставляются с полями **по шапке блока**, не по номеру.
   Пояснение в скобках («Отдел (ОБО, НО, ЮО, HR, ФО)») отбрасывается;
   сравнение идёт от строгого к мягкому, и два одинаково подходящих поля —
   вопрос человеку, а не выбор наугад.
3. **Наши юрлица** предлагаются по колонке исполнителя главного листа;
   написания, отличающиеся кавычками и регистром, сводятся сами.
4. **Статусы** — как написаны; незнакомым смысл назначают сейчас или потом.
5. **Даты окончания** — расторжение или исполнение по статусу и виду; где не
   выводится — «не ясен».
6. **Номера** у разных контрагентов — предупреждение, а не отказ.
7. **Строки вне главного листа** предлагаются к заведению: отбор их сам не
   создаст.
8. **Расхождения** лист ↔ главный лист — по полю; по умолчанию верен главный.
9. **Правила листов** подбираются по строкам листа и показываются с
   покрытием: «340 из 342 строк листа, вот 2 исключения».

Ничего спорного не решается молча. Разобранное лежит в партии загрузки до
«Завести»; файл второй раз не читается.
"""
from __future__ import annotations

import io
import re
import uuid
from collections import Counter as Tally
from dataclasses import dataclass, field
from datetime import date, datetime, timezone
from decimal import Decimal
from typing import Any, Iterable, Sequence

import sqlalchemy as sa
from sqlalchemy.orm import Session

from app.books.layout import norm, squash
from app.finance import history
from app.finance.contracts import setup
from app.finance.contracts.fields import (
    ENTITY,
    SNAPSHOT_FIELDS,
    SYSTEM_KEYS,
    bump,
    deferred_bumps,
    fields_of,
    number_key,
    party_key,
)
from app.finance.contracts.models import Contract, ContractImport, ContractPerson, EntityField, EntityView
from app.finance.contracts.service import (
    RAW_KEY,
    Access,
    Actor,
    Employee,
    Registry,
    _derive,
    _plain,
    _set_field,
    looks_numeric,
    read_money,
)
from app.finance.models import POSITION_STEP, Workspace
from app.finance.service import FinanceError

MAX_ROWS = 20_000
MAX_COLS = 120
#: По сколько договоров «Завести» вставляет в базу за раз.
APPLY_CHUNK = 1_000
#: Шапка — строка, где хотя бы столько ячеек узнаны как поля.
HEADER_MIN_MATCHES = 4
#: Мягкое совпадение шапки: одно начало другого, и доля длины не меньше этой.
LOOSE_MIN_RATIO = 0.6
_PAREN = re.compile(r"\([^)]*\)", re.S)
ROW_NUMBER = "row_number"
#: Поля, различие которых между листами — «расхождение». Снимки оплат не
#: сравниваются: это поле ФО, его и так ведут в одном из листов.
DIFF_FIELDS = (
    "status", "planned_end_at", "people", "number", "signed_at", "department", "type",
    "subject", "amount", "amendments_text", "amendments_summary_text", "end_date", "note", "folder_url",
)
CATEGORICAL = ("type", "subject", "department", "status")
#: Поля, по которым подбирается правило листа, в порядке предпочтения.
#: Статуса здесь нет: лист делит договоры по виду и сторонам, а статус —
#: жизнь договора, и правило «статус = исполнен» увело бы договор из листа,
#: как только его продлили.
RULE_FIELDS = ("subject", "type", "department")


class ImportFailed(FinanceError):
    """Файл нельзя разобрать как реестр."""


# ── Чтение файла ─────────────────────────────────────────────────────────────


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _plain_cell(value: Any) -> Any:
    """Значение ячейки для хранения в партии: даты — ISO, числа — строкой."""
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float, Decimal)):
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value)
    text = str(value).strip()
    return text or None


def _color(fill: Any) -> str:
    try:
        rgb = fill.start_color.rgb if fill is not None and fill.fill_type else None
    except AttributeError:
        return ""
    return rgb[-6:] if isinstance(rgb, str) and len(rgb) >= 6 and rgb not in ("00000000",) else ""


@dataclass
class SheetData:
    name: str
    rows: list[list[Any]]
    widths: dict[int, float] = field(default_factory=dict)
    fills: dict[tuple[int, int], str] = field(default_factory=dict)


def read_workbook(data: bytes) -> list[SheetData]:
    """Все листы книги: значения, ширины колонок и заливка шапок."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise ImportFailed("Не установлен openpyxl") from exc
    try:
        book = load_workbook(io.BytesIO(data), data_only=True)
    except Exception as exc:  # noqa: BLE001 — любой битый файл
        raise ImportFailed("Файл не читается как Excel (.xlsx)") from exc
    out: list[SheetData] = []
    try:
        from openpyxl.utils import column_index_from_string

        for sheet in book.worksheets:
            if sheet.sheet_state != "visible":
                continue
            rows: list[list[Any]] = []
            fills: dict[tuple[int, int], str] = {}
            max_col = min(sheet.max_column or 0, MAX_COLS)
            for r_index, row in enumerate(sheet.iter_rows(max_col=max_col), start=0):
                if r_index >= MAX_ROWS:
                    raise ImportFailed(f"Лист «{sheet.title}»: больше {MAX_ROWS} строк")
                values = []
                for c_index, cell in enumerate(row):
                    values.append(cell.value)
                    if cell.value not in (None, "") and r_index < 400:
                        color = _color(cell.fill)
                        if color:
                            fills[(r_index, c_index)] = color
                rows.append(values)
            widths: dict[int, float] = {}
            for letter, dimension in sheet.column_dimensions.items():
                try:
                    index = column_index_from_string(letter) - 1
                except ValueError:
                    continue
                if dimension.width:
                    widths[index] = round(float(dimension.width), 1)
            out.append(SheetData(sheet.title, rows, widths, fills))
    finally:
        book.close()
    return out


# ── Шапки ────────────────────────────────────────────────────────────────────


def header_core(text: Any) -> str:
    """Шапка без пояснений в скобках: «Отдел (ОБО, НО…)» → «отдел»."""
    return norm(_PAREN.sub(" ", str(text or "")))


@dataclass
class FieldNames:
    key: str
    title: str
    names: set[str]
    squashed: set[str]


def field_names(fields: Sequence[EntityField]) -> list[FieldNames]:
    out = []
    for item in fields:
        names = {norm(name) for name in (item.names or [])} | {norm(item.title)}
        names.discard("")
        out.append(FieldNames(item.key, item.title, names, {squash(name) for name in names}))
    return out


_ROW_NUMBER_NAMES = frozenset({"№", "#", "n", "no", "№ п/п", "п/п"})


class HeaderMatcher:
    """Поиск поля по шапке с разложенным заранее каталогом.

    Разбор проверяет «не шапка ли это» на каждой строке листа, то есть на
    каждой ячейке каждого договора. Прежний поиск на каждую ячейку заново
    нормализовал её текст по разу на каждое поле каталога и на каждый ярус —
    около семидесяти регулярных выражений на ячейку; 20 000 строк разбирались
    пять минут, дольше, чем прокси ждёт ответа. Здесь текст нормализуется один
    раз на ячейку, три строгих яруса — поиск в словаре, а мягкий проверяется,
    только если первые четыре знака совпадают хоть с одним написанием: мягкое
    совпадение — «одно начало другого» длиной от четырёх знаков, без общих
    первых четырёх его не бывает.

    Ответ тот же, что у поиска перебором: ярусы от строгого к мягкому, два
    одинаково подходящих поля — «ambiguous» со списком в порядке каталога.
    """

    def __init__(self, catalog: Sequence[FieldNames]):
        self.catalog = list(catalog)
        self.by_name: dict[str, list[str]] = {}
        self.by_squashed: dict[str, list[str]] = {}
        self.prefixes: set[str] = set()
        for item in self.catalog:
            for name in item.names:
                self.by_name.setdefault(name, []).append(item.key)
            for name in item.squashed:
                self.by_squashed.setdefault(name, []).append(item.key)
                if len(name) >= 4:
                    self.prefixes.add(name[:4])

    def match(self, text: Any) -> tuple[str | None, str, list[str]]:
        full = norm(text)
        if not full:
            return None, "empty", []
        if full in _ROW_NUMBER_NAMES:
            return ROW_NUMBER, "exact", []
        raw = str(text or "")
        # Без скобок `_PAREN` строку не меняет, и ядро совпадает с full.
        core = norm(_PAREN.sub(" ", raw)) if "(" in raw else full
        squashed = squash(core)
        for how, found in (
            ("exact", self.by_name.get(full, [])),
            ("core", self.by_name.get(core, []) if core else []),
            ("squashed", self.by_squashed.get(squashed, []) if squashed else []),
        ):
            if len(found) == 1:
                return found[0], how, []
            if len(found) > 1:
                return None, "ambiguous", list(found)
        if len(squashed) >= 4 and squashed[:4] in self.prefixes:
            found = [item.key for item in self.catalog if _loose(squashed, item.squashed)]
            if len(found) == 1:
                return found[0], "loose", []
            if len(found) > 1:
                return None, "ambiguous", found
        return None, "none", []


def match_header(text: Any, catalog: Sequence[FieldNames]) -> tuple[str | None, str, list[str]]:
    """Поле по шапке: (ключ, как нашлось, кандидаты при двусмысленности).

    Для одной ячейки. Лист целиком разбирается одним `HeaderMatcher`, чтобы
    каталог раскладывался один раз.
    """
    return HeaderMatcher(catalog).match(text)


def _loose(core: str, names: set[str]) -> bool:
    if len(core) < 4:
        return False
    for name in names:
        if len(name) < 4:
            continue
        shorter, longer = sorted((core, name), key=len)
        if longer.startswith(shorter) and len(shorter) / len(longer) >= LOOSE_MIN_RATIO:
            return True
    return False


# ── Блоки ────────────────────────────────────────────────────────────────────


@dataclass
class Column:
    index: int
    header: str
    key: str | None
    how: str
    candidates: list[str] = field(default_factory=list)
    samples: list[str] = field(default_factory=list)
    filled: int = 0
    width: float | None = None


@dataclass
class Block:
    sheet: str
    index: int
    title: str
    header_row: int
    columns: list[Column]
    rows: list[tuple[int, list[Any]]]
    header_fill: str = ""

    @property
    def id(self) -> str:
        return f"{self.sheet}#{self.index}"

    def column_of(self, key: str) -> Column | None:
        return next((column for column in self.columns if column.key == key), None)

    def roles(self) -> dict[str, Any]:
        executor = self.column_of("executor")
        customer = self.column_of("customer")
        out: dict[str, Any] = {}
        if executor is not None:
            out["executor"] = _label(executor.header)
        if customer is not None:
            out["customer"] = _label(customer.header)
        if executor is not None and customer is not None:
            out["order"] = ["executor", "customer"] if executor.index < customer.index else ["customer", "executor"]
        return out


def _label(text: str) -> str:
    return re.sub(r"\s+", " ", str(text or "")).strip()


def _nonempty(row: Sequence[Any]) -> list[tuple[int, Any]]:
    return [(index, value) for index, value in enumerate(row) if value not in (None, "") and str(value).strip()]


def _is_header_row(cells: Sequence[tuple[int, Any]], matcher: HeaderMatcher) -> bool:
    """Узнано ли в строке хотя бы `HEADER_MIN_MATCHES` полей.

    Счёт останавливается, как только ответ ясен: строка договора обычно не
    узнаёт ни одной ячейки, и после того как непроверенных осталось меньше,
    чем не хватает, смотреть их незачем.
    """
    matched = 0
    remaining = len(cells)
    for _, value in cells:
        remaining -= 1
        text = _cell_text(value)
        if text and matcher.match(text)[0] is not None:
            matched += 1
            if matched >= HEADER_MIN_MATCHES:
                return True
        if matched + remaining < HEADER_MIN_MATCHES:
            return False
    return False


def find_blocks(sheet: SheetData, catalog: Sequence[FieldNames]) -> list[Block]:
    """Блоки листа: название → шапка → строки до следующего блока."""
    matcher = HeaderMatcher(catalog)
    headers: list[tuple[int, list[Column]]] = []
    for r_index, row in enumerate(sheet.rows):
        cells = _nonempty(row)
        if len(cells) < HEADER_MIN_MATCHES or not _is_header_row(cells, matcher):
            continue
        columns = []
        for c_index, value in enumerate(row):
            text = _cell_text(value)
            key, how, candidates = matcher.match(text) if text else (None, "empty", [])
            columns.append(Column(c_index, text, key, how, candidates, width=sheet.widths.get(c_index)))
        headers.append((r_index, columns))
    blocks: list[Block] = []
    for number, (h_index, columns) in enumerate(headers):
        previous_end = headers[number - 1][0] if number else -1
        title = ""
        for back in (1, 2, 3):
            probe = h_index - back
            if probe <= previous_end:
                break
            cells = _nonempty(sheet.rows[probe])
            if len(cells) == 1 and isinstance(cells[0][1], str):
                title = _label(cells[0][1])
                break
            if cells:
                break
        next_header = headers[number + 1][0] if number + 1 < len(headers) else len(sheet.rows)
        data: list[tuple[int, list[Any]]] = []
        for r_index in range(h_index + 1, next_header):
            row = sheet.rows[r_index]
            cells = _nonempty(row)
            # Строка договора — хотя бы два узнанных поля. Строка-название
            # следующего блока или пустая строка-отступ сюда не проходят.
            mapped = [
                (c, v)
                for c, v in cells
                if c < len(columns) and columns[c].key not in (None, ROW_NUMBER)
            ]
            if len(mapped) >= 2:
                data.append((r_index, row))
        # Колонки без шапки, но с данными (колонка T «Сводной»).
        width = max((len(row) for _, row in data), default=len(columns))
        for c_index in range(len(columns), width):
            columns.append(Column(c_index, "", None, "empty", width=sheet.widths.get(c_index)))
        for column in columns:
            values = [
                _cell_text(row[column.index]) for _, row in data if column.index < len(row)
            ]
            filled = [value for value in values if value]
            column.filled = len(filled)
            column.samples = list(dict.fromkeys(filled))[:4]
        header_fill = Tally(
            color for (r, c), color in sheet.fills.items() if r == h_index
        ).most_common(1)
        blocks.append(
            Block(
                sheet=sheet.name,
                index=number,
                title=title,
                header_row=h_index + 1,
                columns=[
                    column
                    for column in columns
                    if column.header or column.filled
                ],
                rows=data,
                header_fill=header_fill[0][0] if header_fill else "",
            )
        )
    return blocks


# ── Строки ───────────────────────────────────────────────────────────────────


@dataclass
class Row:
    ref: str
    sheet: str
    block: str
    line: int
    values: dict[str, Any]
    extra: dict[str, Any]

    def to_dict(self) -> dict[str, Any]:
        return {
            "ref": self.ref,
            "sheet": self.sheet,
            "block": self.block,
            "line": self.line,
            "values": self.values,
            "extra": self.extra,
        }


def rows_of(block: Block) -> list[Row]:
    out: list[Row] = []
    for line, raw in block.rows:
        values: dict[str, Any] = {}
        extra: dict[str, Any] = {}
        for column in block.columns:
            if column.index >= len(raw):
                continue
            value = _plain_cell(raw[column.index])
            if value in (None, ""):
                continue
            if column.key == ROW_NUMBER:
                continue
            if column.key is not None:
                values[column.key] = value
            else:
                extra[str(column.index)] = value
        out.append(Row(f"{block.sheet}!{line + 1}", block.sheet, block.id, line + 1, values, extra))
    return out


def row_key(values: dict[str, Any]) -> tuple[str, str, str]:
    return (
        number_key(values.get("number")),
        party_key(values.get("customer")),
        party_key(values.get("executor")),
    )


# ── Разбор: протокол ─────────────────────────────────────────────────────────


def _plural(count: int, one: str, few: str, many: str) -> str:
    tail = count % 100
    if 11 <= tail <= 14:
        return many
    tail = count % 10
    if tail == 1:
        return one
    if 2 <= tail <= 4:
        return few
    return many


def analyze(session: Session, workspace: Workspace, data: bytes, file_name: str) -> dict[str, Any]:
    """Разобрать файл: блоки, колонки, строки. Решений ещё нет."""
    registry = Registry(session, workspace)
    catalog = field_names(registry.fields)
    sheets = read_workbook(data)
    blocks: list[Block] = []
    for sheet in sheets:
        blocks.extend(find_blocks(sheet, catalog))
    if not blocks:
        raise ImportFailed(
            "Не нашли ни одной шапки реестра: ждём строку с колонками вроде «№ Договора», "
            "«Исполнитель», «Заказчик», «Сумма Договора»"
        )
    rows = {block.id: rows_of(block) for block in blocks}
    by_sheet = Tally()
    for block in blocks:
        by_sheet[block.sheet] += len(rows[block.id])
    main_sheet = by_sheet.most_common(1)[0][0]
    return {
        "file": file_name,
        "main_sheet": main_sheet,
        "blocks": [_block_meta(block) for block in blocks],
        "rows": {block_id: [row.to_dict() for row in items] for block_id, items in rows.items()},
    }


def _block_meta(block: Block) -> dict[str, Any]:
    return {
        "id": block.id,
        "sheet": block.sheet,
        "index": block.index,
        "title": block.title,
        "header_row": block.header_row,
        "rows": len(block.rows),
        "header_fill": block.header_fill,
        "roles": block.roles(),
        "columns": [
            {
                "index": column.index,
                "letter": _letter(column.index),
                "header": column.header,
                "key": column.key,
                "how": column.how,
                "candidates": column.candidates,
                "samples": column.samples,
                "filled": column.filled,
                "width": column.width,
            }
            for column in block.columns
        ],
    }


def _letter(index: int) -> str:
    out = ""
    index += 1
    while index:
        index, rest = divmod(index - 1, 26)
        out = chr(65 + rest) + out
    return out


def start(
    session: Session, workspace: Workspace, actor: Actor, data: bytes, file_name: str
) -> ContractImport:
    staged = analyze(session, workspace, data, file_name)
    batch = ContractImport(
        workspace_id=workspace.id,
        file_name=file_name[:200],
        status="preview",
        staged=staged,
        decisions={},
        report={},
        created_by=actor.user_id,
    )
    session.add(batch)
    session.flush()
    batch.report = report(session, workspace, batch)
    session.flush()
    return batch


def get_batch(session: Session, workspace: Workspace, batch_id: uuid.UUID) -> ContractImport:
    batch = session.get(ContractImport, batch_id)
    if batch is None or batch.workspace_id != workspace.id:
        raise FinanceError("Загрузка не найдена")
    return batch


def decide(session: Session, workspace: Workspace, batch_id: uuid.UUID, decisions: dict[str, Any]) -> ContractImport:
    batch = get_batch(session, workspace, batch_id)
    if batch.status != "preview":
        raise FinanceError("Эта загрузка уже заведена или отменена")
    merged = dict(batch.decisions or {})
    for key, value in (decisions or {}).items():
        if isinstance(value, dict) and isinstance(merged.get(key), dict):
            merged[key] = {**merged[key], **value}
        else:
            merged[key] = value
    _check_rule_decisions(merged.get("rules"))
    batch.decisions = merged
    batch.report = report(session, workspace, batch)
    # «Принять как есть» — это то правило, которое человек видел. Записываем
    # его в решение: иначе следующее решение (другой список наших юрлиц)
    # подменило бы принятое правило новым предложением, о котором не спросили.
    rules = dict(merged.get("rules") or {})
    frozen = False
    for item in next(s for s in batch.report["sections"] if s["key"] == "rules")["items"]:
        raw = rules.get(item["block"])
        if isinstance(raw, dict) and raw.get("action") == "accept" and not isinstance(raw.get("filter"), dict):
            rules[item["block"]] = {**raw, "filter": item["filter"]}
            frozen = True
    if frozen:
        batch.decisions = {**merged, "rules": rules}
    session.flush()
    return batch


#: Поля, по которым правило листа проверяется на строках файла, и их условия.
RULE_CONDITION_OPS = {
    **{key: "in" for key in CATEGORICAL},
    "executor_is_own": "is",
    "customer_is_own": "is",
}


def _check_rule_decisions(rules: Any) -> None:
    """Своё правило блока — только из того, что разбор умеет проверить на
    строках файла: вид, предмет, отдел, статус (`in`, значения написаниями) и
    «сторона — наше юрлицо» (`is`). Иначе покрытие в отчёте считалось бы не
    по тому правилу, которое заведётся."""
    if rules in (None, {}):
        return
    if not isinstance(rules, dict):
        raise FinanceError("Решения по правилам листов: ждём {\"<лист>#<блок>\": {...}}")
    for block_id, raw in rules.items():
        if raw is None:
            continue
        if not isinstance(raw, dict) or (raw.get("action") and raw.get("action") not in RULE_ACTIONS):
            raise FinanceError(f"Правило блока {block_id}: ждём action accept, rule или empty")
        rule_filter = raw.get("filter")
        if rule_filter is None:
            continue
        if not isinstance(rule_filter, dict) or not isinstance(rule_filter.get("any", []), list):
            raise FinanceError(f"Правило блока {block_id}: ждём {{\"any\": [{{\"all\": [...]}}]}}")
        for group in rule_filter.get("any") or []:
            conditions = group.get("all") if isinstance(group, dict) else None
            if not isinstance(conditions, list) or not conditions:
                raise FinanceError(f"Правило блока {block_id}: у группы должно быть хотя бы одно условие")
            for condition in conditions:
                field_key = (condition or {}).get("field") if isinstance(condition, dict) else None
                op = RULE_CONDITION_OPS.get(field_key or "")
                if op is None:
                    raise FinanceError(f"Правило блока {block_id}: по полю «{field_key}» при загрузке отбирать нельзя")
                if condition.get("op", "in") != op:
                    raise FinanceError(f"Правило блока {block_id}: для «{field_key}» условие «{op}»")
                if op == "in" and not isinstance(condition.get("value"), list):
                    raise FinanceError(f"Правило блока {block_id}: для «{field_key}» ждём список значений")


# — главное: всё, что протокол знает о строках при текущих решениях —


@dataclass
class Plan:
    """Итог разбора при текущих решениях: что заведётся и что с чем совпало."""

    main_sheet: str
    blocks: list[dict[str, Any]]
    columns: dict[str, dict[str, Any]]
    main_rows: list[dict[str, Any]]
    sheet_rows: dict[str, list[dict[str, Any]]]
    matches: dict[str, str]
    loose_matches: dict[str, str]
    orphans: list[dict[str, Any]]
    existing: dict[str, uuid.UUID]
    own_keys: set[str]
    clusters: dict[str, dict[str, Any]]
    #: Как сведено нестрогое совпадение: `swapped` — те же две стороны, но в
    #: другом порядке; `number_party` — номер и одна общая сторона; `number` —
    #: только номер, по решению человека.
    loose_kinds: dict[str, str] = field(default_factory=dict)
    #: Строка листа → договор главного листа с тем же номером, но без единой
    #: общей стороны. Не сводится, пока человек не скажет «это он».
    number_only: dict[str, str] = field(default_factory=dict)


def _column_decisions(staged: dict[str, Any], decisions: dict[str, Any]) -> dict[str, dict[str, Any]]:
    """Решение по каждой колонке каждого блока: поле, своё поле, примечание, пропуск."""
    chosen = decisions.get("columns") or {}
    out: dict[str, dict[str, Any]] = {}
    for block in staged["blocks"]:
        for column in block["columns"]:
            cid = f"{block['id']}#{column['index']}"
            if cid in chosen:
                out[cid] = chosen[cid]
            elif column["key"]:
                out[cid] = {"action": "field", "field": column["key"]}
            elif column["how"] == "ambiguous":
                out[cid] = {"action": "ask"}
            elif column["header"]:
                out[cid] = {"action": "custom", "title": _label(column["header"]), "type": "text"}
            elif column["filled"]:
                out[cid] = {"action": "ask"}
            else:
                out[cid] = {"action": "skip"}
    return out


def _apply_columns(row: dict[str, Any], block_id: str, columns: dict[str, dict[str, Any]], staged_block: dict[str, Any]) -> dict[str, Any]:
    """Значения строки по решениям о колонках. Своё поле — ключом `custom:<cid>`."""
    values: dict[str, Any] = {}
    notes: list[str] = []
    by_index = {str(column["index"]): column for column in staged_block["columns"]}
    for column in staged_block["columns"]:
        cid = f"{block_id}#{column['index']}"
        decision = columns.get(cid) or {}
        if column["key"]:
            raw = row["values"].get(column["key"])
        else:
            raw = row["extra"].get(str(column["index"]))
        if raw in (None, ""):
            continue
        action = decision.get("action")
        if action == "field" and decision.get("field"):
            values[decision["field"]] = raw
        elif action == "custom":
            values[f"custom:{cid}"] = raw
        elif action == "note":
            label = _label(by_index[str(column["index"])]["header"])
            notes.append(f"{label}: {raw}" if label else str(raw))
    if notes:
        values["note"] = "\n".join([str(values["note"])] + notes) if values.get("note") else "\n".join(notes)
    return values


def plan(session: Session, workspace: Workspace, batch: ContractImport, registry: Registry | None = None) -> Plan:
    staged = batch.staged or {}
    decisions = batch.decisions or {}
    registry = registry or Registry(session, workspace)
    main_sheet = decisions.get("main_sheet") or staged["main_sheet"]
    columns = _column_decisions(staged, decisions)
    blocks = staged["blocks"]
    block_by_id = {block["id"]: block for block in blocks}
    sheet_rows: dict[str, list[dict[str, Any]]] = {}
    main_rows: list[dict[str, Any]] = []
    for block in blocks:
        for row in staged["rows"].get(block["id"], []):
            item = {
                "ref": row["ref"],
                "block": block["id"],
                "sheet": block["sheet"],
                "line": row["line"],
                "values": _apply_columns(row, block["id"], columns, block_by_id[block["id"]]),
            }
            if block["sheet"] == main_sheet:
                main_rows.append(item)
            else:
                sheet_rows.setdefault(block["id"], []).append(item)

    # Совпадения строк листов со строками главного. Строго — номер и обе
    # стороны на своих местах. Нестрого, и такое показывается отдельно:
    # * те же две стороны, переставленные местами (в «Заказчик ГК» и «ФИН.
    #   ПОМОЩЬ» реестра BBC все шесть нестрогих совпадений — такие);
    # * номер, однозначный с обеих сторон, и хотя бы одна общая сторона, не
    #   наше юрлицо (другое написание второй стороны или её замена). Наше
    #   юрлицо стоит почти в каждом договоре — общим оно доводом не считается.
    # Один номер без общей внешней стороны — не совпадение: номер у разных
    # клиентов повторяется, это данные. Такая строка раньше приклеивалась к
    # чужому договору: сама не заводилась, а правило листа подбиралось по
    # признакам чужого договора и тянуло в лист тысячи лишних. Теперь она
    # заводится отдельно, а человек может сказать «это он»
    # (`decisions["loose"][ref] = "same"`).
    main_clusters = _party_clusters(main_rows, registry, decisions, main_refs={row["ref"] for row in main_rows})
    own_for_matching = _own_keys(main_clusters, decisions) | set(
        (decisions.get("entities") or {}).get("own") or []
    )
    main_by_key: dict[tuple[str, str, str], list[str]] = {}
    main_by_number: dict[str, list[str]] = {}
    main_key_of: dict[str, tuple[str, str, str]] = {}
    for item in main_rows:
        key = row_key(item["values"])
        main_key_of[item["ref"]] = key
        main_by_key.setdefault(key, []).append(item["ref"])
        if key[0]:
            main_by_number.setdefault(key[0], []).append(item["ref"])
    matches: dict[str, str] = {}
    loose: dict[str, str] = {}
    loose_kinds: dict[str, str] = {}
    number_only: dict[str, str] = {}
    orphans: list[dict[str, Any]] = []
    sheet_number_counts: dict[str, Tally] = {}
    for block_id, items in sheet_rows.items():
        tally = Tally(number_key(item["values"].get("number")) for item in items)
        sheet_number_counts[block_id] = tally
    choices = decisions.get("loose") or {}
    separate = {ref for ref, choice in choices.items() if choice == "separate"}
    same = {ref for ref, choice in choices.items() if choice == "same"}
    for block_id, items in sheet_rows.items():
        for item in items:
            ref = item["ref"]
            key = row_key(item["values"])
            exact = main_by_key.get(key) or []
            if len(exact) >= 1:
                matches[ref] = exact[0]
                continue
            if ref in separate:
                orphans.append(item)
                continue
            swapped = main_by_key.get((key[0], key[2], key[1])) or [] if key[0] else []
            if swapped:
                loose[ref] = swapped[0]
                loose_kinds[ref] = "swapped"
                continue
            by_number = main_by_number.get(key[0]) if key[0] else None
            if by_number and len(by_number) == 1 and sheet_number_counts[block_id][key[0]] == 1:
                main_key = main_key_of[by_number[0]]
                common = ({key[1], key[2]} & {main_key[1], main_key[2]}) - {""} - own_for_matching
                if common or ref in same:
                    loose[ref] = by_number[0]
                    loose_kinds[ref] = "number_party" if common else "number"
                    continue
                number_only[ref] = by_number[0]
            orphans.append(item)

    # Существующие договоры компании (повторная загрузка): совпадение по тому
    # же ключу превращает строку файла из «завести» в «сверить».
    existing: dict[str, uuid.UUID] = {}
    live = session.scalars(
        sa.select(Contract).where(Contract.workspace_id == workspace.id, Contract.deleted_at.is_(None))
    ).all()
    if live:
        by_existing_key: dict[tuple[str, str, str], uuid.UUID] = {}
        for contract in live:
            customer = registry.parties.get(contract.customer_id) if contract.customer_id else None
            executor = registry.parties.get(contract.executor_id) if contract.executor_id else None
            by_existing_key[
                (
                    contract.number_key,
                    party_key(customer.name) if customer else "",
                    party_key(executor.name) if executor else "",
                )
            ] = contract.id
        for item in main_rows + orphans:
            found = by_existing_key.get(row_key(item["values"]))
            if found is not None:
                existing[item["ref"]] = found

    clusters = _party_clusters(main_rows + orphans, registry, decisions, main_refs={row["ref"] for row in main_rows})
    own_keys = _own_keys(clusters, decisions)
    return Plan(
        main_sheet, blocks, columns, main_rows, sheet_rows, matches, loose, orphans, existing, own_keys, clusters,
        loose_kinds=loose_kinds, number_only=number_only,
    )


def _party_clusters(
    rows: Iterable[dict[str, Any]],
    registry: Registry,
    decisions: dict[str, Any],
    *,
    main_refs: set[str] | frozenset[str] = frozenset(),
) -> dict[str, dict[str, Any]]:
    """Стороны файла, сведённые по ключу: написания, где стоят, сколько раз.

    `main_executor` / `main_customer` — только по главному листу: он ведётся
    от лица группы, и исполнитель там почти всегда наше ТОО. В других листах
    исполнителем бывает внешняя сторона (продавец, у которого мы покупаем), и
    считать её там значило бы предложить нашим того, у кого мы заказчик.
    """
    merges = decisions.get("merges") or []
    alias: dict[str, str] = {}
    for group in merges:
        keys = [key for key in group if key]
        for key in keys[1:]:
            alias[key] = keys[0]
    clusters: dict[str, dict[str, Any]] = {}
    for row in rows:
        for slot in ("executor", "customer"):
            name = row["values"].get(slot)
            if not name:
                continue
            key = party_key(name)
            key = alias.get(key, key)
            cluster = clusters.setdefault(
                key,
                {
                    "key": key, "names": Tally(), "executor": 0, "customer": 0,
                    "main_executor": 0, "main_customer": 0, "known": None, "own": False,
                },
            )
            cluster["names"][str(name).strip()] += 1
            cluster[slot] += 1
            if row["ref"] in main_refs:
                cluster[f"main_{slot}"] += 1
    for cluster in clusters.values():
        name = cluster["names"].most_common(1)[0][0]
        resolved = registry.resolve_party(name, slot="executor", create=False)
        if resolved.party is not None:
            cluster["known"] = str(resolved.party.id)
            cluster["own"] = registry.is_own(resolved.party.id)
    return clusters


def _own_keys(clusters: dict[str, dict[str, Any]], decisions: dict[str, Any]) -> set[str]:
    chosen = (decisions.get("entities") or {}).get("own")
    if chosen is not None:
        return {key for key in chosen if key in clusters}
    return {key for key, cluster in clusters.items() if cluster["own"] or _proposed_own(cluster)}


def _proposed_own(cluster: dict[str, Any]) -> bool:
    """Предложить нашим: исполнитель хотя бы в двух строках главного листа и
    исполнителем бывает чаще, чем заказчиком.

    Это предложение, а не решение: список подтверждает человек, и число строк
    стоит рядом с каждым именем. Частное лицо, которое и даёт заём, и берёт,
    сюда не попадёт; ТОО группы с одной строкой человек отметит сам.
    """
    return cluster["main_executor"] >= 2 and cluster["main_executor"] > cluster["main_customer"]


def _similar_clusters(clusters: dict[str, dict[str, Any]]) -> list[list[str]]:
    """Пары написаний, которые похожи, но по ключу не совпали: «ТОО Альфа» и «Альфа»."""
    stripped: dict[str, list[str]] = {}
    for key in clusters:
        bare = re.sub(r"^(тоо|ип|ао|too|llp)", "", key)
        if len(bare) >= 4:
            stripped.setdefault(bare, []).append(key)
    return [sorted(keys) for keys in stripped.values() if len(keys) > 1]


# ── Отчёт: девять пунктов ────────────────────────────────────────────────────


def report(session: Session, workspace: Workspace, batch: ContractImport) -> dict[str, Any]:
    return _report_and_plan(session, workspace, batch)[0]


def _report_and_plan(
    session: Session, workspace: Workspace, batch: ContractImport, registry: Registry | None = None
) -> tuple[dict[str, Any], Plan]:
    """Отчёт и план, по которому он посчитан: «Завести» берёт план отсюда, а
    не считает его второй раз."""
    registry = registry or Registry(session, workspace)
    current_plan = plan(session, workspace, batch, registry)
    staged = batch.staged or {}
    decisions = batch.decisions or {}
    sections = [
        _section_blocks(current_plan, staged),
        _section_columns(current_plan, registry),
        _section_entities(current_plan, decisions),
        _section_statuses(current_plan, registry, decisions),
        _section_end_dates(current_plan, registry, decisions),
        _section_numbers(current_plan),
        _section_orphans(current_plan, decisions),
        _section_diffs(current_plan, registry, decisions),
        _section_rules(current_plan, registry, decisions),
    ]
    blocking = [section["key"] for section in sections if section["blocking"] and not section["done"]]
    creating = len([row for row in current_plan.main_rows if row["ref"] not in current_plan.existing]) + len(
        [
            row
            for row in current_plan.orphans
            if (decisions.get("orphans") or {}).get(row["ref"], True) and row["ref"] not in current_plan.existing
        ]
    )
    return {
        "file": staged.get("file", ""),
        "main_sheet": current_plan.main_sheet,
        "sheets": sorted({block["sheet"] for block in current_plan.blocks}, key=lambda name: [b["sheet"] for b in current_plan.blocks].index(name)),
        "sections": sections,
        "blocking": blocking,
        "totals": {
            "create": creating,
            "update": len(current_plan.existing),
            "main_rows": len(current_plan.main_rows),
        },
    }, current_plan


def _section(key: str, title: str, *, blocking: bool, done: bool, summary: str, items: list[Any], **extra: Any) -> dict[str, Any]:
    return {"key": key, "title": title, "blocking": blocking, "done": done, "summary": summary, "items": items, **extra}


def _section_blocks(current: Plan, staged: dict[str, Any]) -> dict[str, Any]:
    items = []
    for block in current.blocks:
        items.append(
            {
                "id": block["id"],
                "sheet": block["sheet"],
                "title": block["title"],
                "rows": block["rows"],
                "header_row": block["header_row"],
                "roles": block["roles"],
                "main": block["sheet"] == current.main_sheet,
            }
        )
    count = len(items)
    return _section(
        "blocks",
        "Листы и блоки",
        blocking=False,
        done=True,
        summary=f"{count} {_plural(count, 'блок', 'блока', 'блоков')}; главный лист — «{current.main_sheet}»",
        items=items,
    )


def _section_columns(current: Plan, registry: Registry) -> dict[str, Any]:
    items = []
    open_questions = 0
    for block in current.blocks:
        columns = []
        for column in block["columns"]:
            cid = f"{block['id']}#{column['index']}"
            decision = current.columns.get(cid, {})
            if decision.get("action") == "ask":
                open_questions += 1
            columns.append({**column, "id": cid, "decision": decision})
        items.append({"block": block["id"], "sheet": block["sheet"], "title": block["title"], "roles": block["roles"], "columns": columns})
    return _section(
        "columns",
        "Колонки",
        blocking=True,
        done=open_questions == 0,
        summary=(
            "Все колонки легли на поля"
            if open_questions == 0
            else f"{open_questions} {_plural(open_questions, 'колонка ждёт', 'колонки ждут', 'колонок ждут')} решения"
        ),
        items=items,
    )


def _section_entities(current: Plan, decisions: dict[str, Any]) -> dict[str, Any]:
    items = []
    for key, cluster in sorted(current.clusters.items(), key=lambda pair: -(pair[1]["executor"] + pair[1]["customer"])):
        if not (cluster["executor"] >= 1 or cluster["own"]):
            continue
        items.append(
            {
                "key": key,
                "names": [name for name, _ in cluster["names"].most_common()],
                "executor": cluster["executor"],
                "customer": cluster["customer"],
                "own": key in current.own_keys,
                "known": cluster["known"],
            }
        )
    similar = [
        {"keys": group, "names": [current.clusters[key]["names"].most_common(1)[0][0] for key in group]}
        for group in _similar_clusters(current.clusters)
    ]
    confirmed = bool((decisions.get("entities") or {}).get("confirmed"))
    own_count = len(current.own_keys)
    spellings = sum(len(current.clusters[key]["names"]) for key in current.own_keys)
    return _section(
        "entities",
        "Наши юрлица",
        blocking=True,
        done=confirmed,
        summary=(
            f"{spellings} {_plural(spellings, 'написание', 'написания', 'написаний')} → {own_count} "
            f"{_plural(own_count, 'юрлицо', 'юрлица', 'юрлиц')}"
            + ("" if confirmed else " — подтвердите список")
        ),
        items=items,
        similar=similar,
    )


def _values_of(rows: Iterable[dict[str, Any]], key: str) -> Tally:
    return Tally(_label(row["values"].get(key)) for row in rows if row["values"].get(key))


def _section_statuses(current: Plan, registry: Registry, decisions: dict[str, Any]) -> dict[str, Any]:
    rows = current.main_rows + current.orphans
    items = []
    unknown = 0
    for field_key in ("status", "type", "department"):
        for value, count in _values_of(rows, field_key).most_common():
            if field_key == "department":
                existing = registry.resolve_department(value, create=False)
                meaning: dict[str, Any] = {}
                known = existing is not None
                odd = bool(re.search(r"[,;\n]", value))
            else:
                existing = registry.resolve_value(field_key, value, create=False)
                meaning = dict(existing.meaning or {}) if existing is not None else {}
                chosen = ((decisions.get("statuses") or {}).get(f"{field_key}:{value}")) or {}
                meaning = {**meaning, **chosen}
                known = bool(meaning)
                odd = count == 1 and field_key == "type"
            if not known and field_key == "status":
                unknown += 1
            items.append(
                {"field": field_key, "value": value, "count": count, "known": known, "meaning": meaning, "odd": odd}
            )
    return _section(
        "statuses",
        "Статусы и списки",
        blocking=False,
        done=True,
        summary=(
            "Все статусы знакомы"
            if not unknown
            else f"{unknown} {_plural(unknown, 'статус', 'статуса', 'статусов')} без смысла — сохранятся как написаны"
        ),
        items=items,
    )


def _end_kind_of(values: dict[str, Any], registry: Registry, decisions: dict[str, Any]) -> str:
    status = registry.resolve_value("status", values.get("status"), create=False) if values.get("status") else None
    chosen = ((decisions.get("statuses") or {}).get(f"status:{_label(values.get('status'))}")) or {}
    phase = chosen.get("phase") or ((status.meaning or {}).get("phase") if status is not None else None)
    kind_value = registry.resolve_value("type", values.get("type"), create=False) if values.get("type") else None
    billing = (kind_value.meaning or {}).get("billing") if kind_value is not None else None
    if phase == "terminated":
        return "terminated"
    if phase == "fulfilled" and billing == "total":
        return "fulfilled"
    return "unknown"


def _section_end_dates(current: Plan, registry: Registry, decisions: dict[str, Any]) -> dict[str, Any]:
    chosen = decisions.get("end_dates") or {}
    tally = Tally()
    unclear = []
    for row in current.main_rows + current.orphans:
        if not row["values"].get("end_date"):
            continue
        kind = chosen.get(row["ref"]) or _end_kind_of(row["values"], registry, decisions)
        tally[kind] += 1
        if kind == "unknown":
            unclear.append(
                {
                    "ref": row["ref"],
                    "number": row["values"].get("number"),
                    "customer": row["values"].get("customer"),
                    "status": row["values"].get("status"),
                    "type": row["values"].get("type"),
                    "end_date": row["values"].get("end_date"),
                }
            )
    return _section(
        "end_dates",
        "Даты окончания",
        blocking=False,
        done=True,
        summary=(
            f"расторжение {tally['terminated']}, исполнение {tally['fulfilled']}, смысл не ясен {tally['unknown']}"
        ),
        items=unclear,
        counts=dict(tally),
    )


def _section_numbers(current: Plan) -> dict[str, Any]:
    by_number: dict[str, dict[str, set[str]]] = {}
    for row in current.main_rows + current.orphans:
        key = number_key(row["values"].get("number"))
        if not key:
            continue
        entry = by_number.setdefault(key, {"number": set(), "parties": set(), "refs": set()})
        entry["number"].add(str(row["values"].get("number")).strip())
        entry["parties"].add(party_key(row["values"].get("customer")) + "|" + party_key(row["values"].get("executor")))
        entry["refs"].add(row["ref"])
    items = [
        {"number": sorted(entry["number"])[0], "contracts": len(entry["refs"]), "refs": sorted(entry["refs"])}
        for entry in by_number.values()
        if len(entry["parties"]) > 1
    ]
    count = len(items)
    return _section(
        "numbers",
        "Номера у разных контрагентов",
        blocking=False,
        done=True,
        summary=(
            "Повторов нет"
            if not count
            else f"{count} {_plural(count, 'номер стоит', 'номера стоят', 'номеров стоят')} у разных контрагентов — "
            "договоры заведутся, у каждого будет замечание"
        ),
        items=items,
    )


def _section_orphans(current: Plan, decisions: dict[str, Any]) -> dict[str, Any]:
    chosen = decisions.get("orphans") or {}
    items = []
    for row in current.orphans:
        items.append(
            {
                "ref": row["ref"],
                "block": row["block"],
                "number": row["values"].get("number"),
                "executor": row["values"].get("executor"),
                "customer": row["values"].get("customer"),
                "create": chosen.get(row["ref"], True),
                "existing": str(current.existing[row["ref"]]) if row["ref"] in current.existing else None,
            }
        )
    main_by_ref = {row["ref"]: row for row in current.main_rows}
    sheet_by_ref = {row["ref"]: row for rows in current.sheet_rows.values() for row in rows}

    def pair(ref: str, main_ref: str, **extra: Any) -> dict[str, Any]:
        sheet_values, main_values = sheet_by_ref[ref]["values"], main_by_ref[main_ref]["values"]
        return {
            "ref": ref,
            "main": main_ref,
            "number": sheet_values.get("number"),
            "sheet_customer": sheet_values.get("customer"),
            "sheet_executor": sheet_values.get("executor"),
            "main_customer": main_values.get("customer"),
            "main_executor": main_values.get("executor"),
            **extra,
        }

    # Сведены нестрого — показываются, чтобы человек мог развести («separate»).
    loose = [
        pair(ref, main_ref, kind=current.loose_kinds.get(ref, "number_party"))
        for ref, main_ref in current.loose_matches.items()
    ]
    # Тот же номер, но ни одной общей стороны — заведутся отдельно, пока
    # человек не скажет «это он» («same»).
    number_only = [pair(ref, main_ref) for ref, main_ref in current.number_only.items()]
    count = len(items)
    return _section(
        "orphans",
        f"Строки, которых нет в «{current.main_sheet}»",
        blocking=False,
        done=True,
        summary=(
            f"{count} {_plural(count, 'строка', 'строки', 'строк')} — заведутся как договоры"
            if count
            else "Все строки листов есть в главном листе"
        )
        + (f"; {len(loose)} сведены нестрого" if loose else "")
        + (f"; {len(number_only)} совпали только номером — заведутся отдельно" if number_only else ""),
        items=items,
        loose=loose,
        number_only=number_only,
    )


def _normalize_for_diff(key: str, value: Any) -> str:
    if value in (None, ""):
        return ""
    if key in ("amount",):
        amount, terms = read_money(value, field="Сумма") if looks_numeric(str(value)) else (None, str(value))
        return format(amount, "f") if amount is not None else squash(terms)
    if key in ("executor", "customer"):
        return party_key(value)
    return squash(value)


def _section_diffs(current: Plan, registry: Registry, decisions: dict[str, Any]) -> dict[str, Any]:
    chosen = decisions.get("diffs") or {}
    main_by_ref = {row["ref"]: row for row in current.main_rows}
    items = []
    by_field = Tally()
    for block_id, rows in current.sheet_rows.items():
        for row in rows:
            main_ref = current.matches.get(row["ref"]) or current.loose_matches.get(row["ref"])
            if not main_ref:
                continue
            main = main_by_ref[main_ref]
            for key in DIFF_FIELDS:
                a, b = main["values"].get(key), row["values"].get(key)
                try:
                    differ = _normalize_for_diff(key, a) != _normalize_for_diff(key, b)
                except FinanceError:
                    differ = str(a or "") != str(b or "")
                if not differ or not b:
                    continue
                diff_id = f"{main_ref}|{row['ref']}|{key}"
                by_field[key] += 1
                items.append(
                    {
                        "id": diff_id,
                        "main_ref": main_ref,
                        "sheet_ref": row["ref"],
                        "field": key,
                        "title": registry.field_by_key[key].title if key in registry.field_by_key else key,
                        "main": a,
                        "sheet": b,
                        "number": main["values"].get("number"),
                        "customer": main["values"].get("customer"),
                        "take": chosen.get(diff_id, "main"),
                    }
                )
    count = len(items)
    detail = ", ".join(
        f"{registry.field_by_key[key].title if key in registry.field_by_key else key} {n}"
        for key, n in by_field.most_common(6)
    )
    return _section(
        "diffs",
        "Расхождения листов",
        blocking=False,
        done=True,
        summary=(
            "Листы совпадают с главным"
            if not count
            else f"{count} {_plural(count, 'расхождение', 'расхождения', 'расхождений')} ({detail}); "
            f"по умолчанию верен «{current.main_sheet}»"
        ),
        items=items,
    )


# — правила листов —


def _facts_for_rule(row: dict[str, Any], own_keys: set[str]) -> dict[str, Any]:
    values = row["values"]
    facts = {key: _label(values.get(key)) for key in CATEGORICAL if values.get(key)}
    executor_own = party_key(values.get("executor")) in own_keys if values.get("executor") else False
    customer_own = party_key(values.get("customer")) in own_keys if values.get("customer") else False
    facts["executor_is_own"] = executor_own
    facts["customer_is_own"] = customer_own
    return facts


def _rule_matches(rule: list[list[tuple[str, Any]]], facts: dict[str, Any]) -> bool:
    for group in rule:
        ok = True
        for field_key, wanted in group:
            if isinstance(wanted, bool):
                if bool(facts.get(field_key)) is not wanted:
                    ok = False
                    break
            elif facts.get(field_key) not in wanted:
                ok = False
                break
        if ok:
            return True
    return False


def _best_group(target: list[dict[str, Any]], universe: list[dict[str, Any]]) -> list[tuple[str, Any]]:
    """Лучшая группа условий «и» для строк блока: сперва одно поле, потом второе.

    Кандидат — «поле ∈ значения этих строк» (если значений немного) или
    виртуальный признак стороны. Выигрывает тот, что ловит все строки блока и
    меньше всего чужих; второе условие добавляется, если убирает чужие, не
    теряя своих.
    """
    candidates: list[tuple[str, Any]] = []
    for key in RULE_FIELDS:
        values = {facts.get(key) for facts in target if facts.get(key)}
        if values and len(values) <= 5 and all(facts.get(key) for facts in target):
            candidates.append((key, sorted(values)))
    own_sides: list[tuple[str, Any]] = []
    for key in ("executor_is_own", "customer_is_own"):
        flags = {bool(facts.get(key)) for facts in target}
        if len(flags) == 1:
            flag = flags.pop()
            candidates.append((key, flag))
            if flag:
                own_sides.append((key, True))

    def score(group: list[tuple[str, Any]]) -> tuple[int, int, int]:
        caught = sum(1 for facts in universe if _rule_matches([group], facts))
        own = sum(1 for facts in target if _rule_matches([group], facts))
        # При равном покрытии — правило короче: «вид ∈ {Аренда}» надёжнее двух
        # написаний предмета, новая аренда с третьим написанием не выпадет.
        listed = sum(len(wanted) if isinstance(wanted, list) else 0 for _, wanted in group)
        return own, -caught, -listed

    best: list[tuple[str, Any]] = []
    best_score = (0, -len(universe), 0)
    for candidate in candidates:
        current = score([candidate])
        if current > best_score:
            best, best_score = [candidate], current
    if best:
        for candidate in candidates:
            if any(candidate[0] == item[0] for item in best):
                continue
            trial = best + [candidate]
            current = score(trial)
            if current[0] >= best_score[0] and current[1] > best_score[1]:
                best, best_score = trial, (current[0], current[1], best_score[2])
        # Сторона группы, верная для всех строк листа, входит в правило, даже
        # если лишних договоров сейчас не убирает: «Исполнитель ГК» — это лист,
        # где исполнитель наш, и договор с чужим исполнителем туда не должен
        # попасть завтра.
        for side in own_sides:
            if not any(side[0] == item[0] for item in best):
                trial = best + [side]
                if score(trial)[0] >= best_score[0]:
                    best, best_score = trial, score(trial)
    return best


def suggest_rule(target_rows: list[dict[str, Any]], universe_rows: list[dict[str, Any]], own_keys: set[str]) -> list[list[tuple[str, Any]]]:
    target = [_facts_for_rule(row, own_keys) for row in target_rows]
    universe = [_facts_for_rule(row, own_keys) for row in universe_rows]
    if not target:
        return []
    groups: list[list[tuple[str, Any]]] = []
    remaining = target
    for _ in range(3):
        group = _best_group(remaining, universe)
        if not group:
            break
        groups.append(group)
        remaining = [facts for facts in remaining if not _rule_matches([group], facts)]
        if not remaining:
            break
    return groups


def _rule_to_filter(rule: list[list[tuple[str, Any]]]) -> dict[str, Any]:
    """Правило в тексте (значения — написания) → фильтр листа для отчёта."""
    return {
        "any": [
            {
                "all": [
                    {"field": key, "op": "is", "value": wanted}
                    if isinstance(wanted, bool)
                    else {"field": key, "op": "in", "value": list(wanted)}
                    for key, wanted in group
                ]
            }
            for group in rule
        ]
    }


def _filter_to_rule(rule_filter: dict[str, Any]) -> list[list[tuple[str, Any]]]:
    out = []
    for group in (rule_filter or {}).get("any") or []:
        items = []
        for condition in group.get("all") or []:
            if condition.get("op") == "is":
                items.append((condition["field"], bool(condition.get("value"))))
            else:
                items.append((condition["field"], [_label(v) for v in condition.get("value") or []]))
        out.append(items)
    return out


def _block_target(current: Plan, block_id: str, decisions: dict[str, Any]) -> list[dict[str, Any]]:
    """Строки, которые будут договорами и стоят в этом блоке."""
    main_by_ref = {row["ref"]: row for row in current.main_rows}
    orphan_by_ref = {row["ref"]: row for row in current.orphans}
    out = []
    for row in current.sheet_rows.get(block_id, []):
        ref = current.matches.get(row["ref"]) or current.loose_matches.get(row["ref"])
        if ref:
            out.append(main_by_ref[ref])
        elif row["ref"] in orphan_by_ref and (decisions.get("orphans") or {}).get(row["ref"], True):
            out.append(orphan_by_ref[row["ref"]])
    return out


def _universe(current: Plan, decisions: dict[str, Any]) -> list[dict[str, Any]]:
    chosen = decisions.get("orphans") or {}
    return current.main_rows + [row for row in current.orphans if chosen.get(row["ref"], True)]


#: Когда предложенное правило листа не принимается без человека: лишних
#: договоров (правило добавит их в лист, а в листе файла их нет) или
#: недостающих больше пяти строк или больше десятой части листа.
#:
#: Пять строк — столько человек проверит глазами по образцу в отчёте (там до
#: десяти лишних и до двадцати недостающих). Десятая часть — граница, за
#: которой лист уже не «тот же, что в Excel»: маленький лист упирается в неё
#: раньше (один лишний на восемь строк), большой — в пять строк. На реестре
#: BBC правило пропускается без вопроса у «Исполнитель ГК» (5 лишних на 341)
#: и «ФИН. ПОМОЩЬ» (2 на 21), а у «Заказчик ГК» (47 лишних на 16) — нет:
#: там лист после загрузки был бы вчетверо больше, чем в файле.
RULE_STRAY_ROWS = 5
RULE_STRAY_SHARE = 0.10
RULE_ACTIONS = ("accept", "rule", "empty")


def rule_limit(in_sheet: int) -> int:
    """Сколько лишних или недостающих строк правило может принести без вопроса."""
    return min(RULE_STRAY_ROWS, int(in_sheet * RULE_STRAY_SHARE))


def _rule_decision(raw: Any) -> tuple[str | None, dict[str, Any] | None]:
    """Решение человека по правилу блока: (действие, правило-написаниями).

    `{"filter": …}` без действия — прежний формат «своё правило».
    """
    if not isinstance(raw, dict):
        return None, None
    action = raw.get("action") or ("rule" if raw.get("filter") else None)
    if action not in RULE_ACTIONS:
        return None, None
    rule_filter = raw.get("filter") if isinstance(raw.get("filter"), dict) else None
    if action == "rule" and not (rule_filter or {}).get("any"):
        return "empty", None
    return action, rule_filter


def _section_rules(current: Plan, registry: Registry, decisions: dict[str, Any]) -> dict[str, Any]:
    """Правило каждого блока других листов, его покрытие и нужен ли человек.

    Решение по блоку — `decisions["rules"]["<лист>#<блок>"]`:
    * `{"action": "accept"}` — принять предложенное как есть; на `decide`
      предложенное правило записывается в решение (`filter`), чтобы следующие
      решения его не подменили;
    * `{"action": "rule", "filter": {"any": [...]}}` — своё правило;
    * `{"action": "empty"}` — блок без правила: договоров в нём не будет,
      пока правило не задано в настройке листа.
    """
    chosen = decisions.get("rules") or {}
    universe = _universe(current, decisions)
    universe_facts = {row["ref"]: _facts_for_rule(row, current.own_keys) for row in universe}
    by_ref = {row["ref"]: row for row in universe}
    items = []
    pending: list[str] = []
    for block in current.blocks:
        if block["sheet"] == current.main_sheet:
            continue
        target = _block_target(current, block["id"], decisions)
        action, rule_filter = _rule_decision(chosen.get(block["id"]))
        if action == "empty":
            rule, source = [], "empty"
        elif action == "rule":
            rule, source = _filter_to_rule(rule_filter), "manual"
        elif action == "accept" and rule_filter is not None:
            rule, source = _filter_to_rule(rule_filter), "accepted"
        else:
            rule = suggest_rule(target, universe, current.own_keys)
            source = "accepted" if action == "accept" else "suggested"
        target_refs = {row["ref"] for row in target}
        caught = {ref for ref, facts in universe_facts.items() if rule and _rule_matches(rule, facts)}
        missing = sorted(target_refs - caught)
        extra = sorted(caught - target_refs)
        limit = rule_limit(len(target_refs))
        reason = ""
        if target_refs and source == "suggested":
            if not rule:
                reason = "правило не подобралось"
            elif len(extra) > limit or len(missing) > limit:
                parts = []
                if len(extra) > limit:
                    parts.append(f"лишних {len(extra)}")
                if len(missing) > limit:
                    parts.append(f"не попадут {len(missing)}")
                reason = f"{', '.join(parts)} при {len(target_refs)} строках листа — допустимо не больше {limit}"
        needs_decision = bool(reason)
        if needs_decision:
            pending.append(block["id"])
        items.append(
            {
                "block": block["id"],
                "sheet": block["sheet"],
                "title": block["title"],
                "filter": _rule_to_filter(rule),
                "sentence": (
                    "Блок без правила — договоров в нём нет, пока правило не задано в настройке листа"
                    if source == "empty"
                    else _sentence(rule, registry)
                ),
                "source": source,
                "decision": action,
                "needs_decision": needs_decision,
                "reason": reason,
                "limit": limit,
                "in_sheet": len(target_refs),
                "caught": len(caught & target_refs),
                "extra": len(extra),
                "missing_count": len(missing),
                "missing": [_brief(by_ref[ref]) for ref in missing[:20]],
                "extra_sample": [_brief(by_ref[ref]) for ref in extra[:10]],
            }
        )
    overview = "; ".join(
        f"{item['sheet']}{(' / ' + item['title']) if item['title'] else ''}: {item['caught']} из {item['in_sheet']}"
        + (f", лишних {item['extra']}" if item["extra"] else "")
        for item in items[:6]
    )
    count = len(pending)
    return _section(
        "rules",
        "Правила листов",
        blocking=True,
        done=not pending,
        summary=(
            f"{count} {_plural(count, 'правило ждёт', 'правила ждут', 'правил ждут')} решения — {overview}"
            if pending
            else (overview or "Других листов нет")
        ),
        items=items,
        pending=pending,
    )


def _brief(row: dict[str, Any]) -> dict[str, Any]:
    values = row["values"]
    return {
        "ref": row["ref"],
        "number": values.get("number"),
        "customer": values.get("customer"),
        "executor": values.get("executor"),
        "type": values.get("type"),
        "subject": values.get("subject"),
    }


_FIELD_WORDS = {
    "type": "вид",
    "subject": "предмет",
    "department": "отдел",
    "status": "статус",
    "executor_is_own": "исполнитель — наше юрлицо",
    "customer_is_own": "заказчик — наше юрлицо",
}


def _sentence(rule: list[list[tuple[str, Any]]], registry: Registry) -> str:
    if not rule:
        return "Правило не подобралось — задайте его в настройке листа"
    groups = []
    for group in rule:
        parts = []
        for key, wanted in group:
            if isinstance(wanted, bool):
                word = _FIELD_WORDS.get(key, key)
                parts.append(word if wanted else f"не {word}")
            else:
                parts.append(f"{_FIELD_WORDS.get(key, key)} входит в {', '.join(wanted)}")
        groups.append(" и ".join(parts))
    return "Показывать договоры, где " + ", или ".join(groups)


# ── Заведение ────────────────────────────────────────────────────────────────


def apply(session: Session, workspace: Workspace, access: Access, actor: Actor, batch_id: uuid.UUID) -> dict[str, Any]:
    """Завести партию. Номер схемы двигается один раз — на выходе.

    Заведение юрлиц, значений списков, отделов, полей и листов каждый раз
    двигало номер схемы в начале транзакции и держало его блокировку до
    коммита: правка коллеги, заводящая новое значение списка, ждала всю
    загрузку (см. `fields.deferred_bumps`).
    """
    with deferred_bumps(session, workspace.id, "schema"):
        return _apply(session, workspace, access, actor, batch_id)


def _apply(session: Session, workspace: Workspace, access: Access, actor: Actor, batch_id: uuid.UUID) -> dict[str, Any]:
    """Завести договоры, листы, наши юрлица и свои поля по решениям.

    Номер изменения реестра (`seq`) — один на всю партию и берётся последним,
    перед коммитом. Счётчик держит блокировку строки до конца транзакции
    (так опрос видит номера строго по порядку), и прежде он брался до
    заведения строк: пока заводились 20 000 договоров — две минуты с лишним —
    любая правка в этой компании ждала. Теперь договоры вставляются пачками по
    `APPLY_CHUNK` с `seq = 0` (до коммита их никто не видит), и только потом
    берётся номер и одним UPDATE проставляется всей партии.
    """
    if not access.setup:
        raise PermissionError("Загружать реестр может владелец или администратор")
    batch = get_batch(session, workspace, batch_id)
    if batch.status != "preview":
        raise FinanceError("Эта загрузка уже заведена или отменена")
    registry = Registry(session, workspace)
    batch.report, current = _report_and_plan(session, workspace, batch, registry)
    if batch.report["blocking"]:
        titles = [s["title"] for s in batch.report["sections"] if s["key"] in batch.report["blocking"]]
        raise FinanceError("Сначала решите: " + ", ".join(titles))
    decisions = batch.decisions or {}

    # 1. Наши юрлица — все написания каждого становятся псевдонимами.
    own_party: dict[str, uuid.UUID] = {}
    for key in current.own_keys:
        cluster = current.clusters[key]
        names = [name for name, _ in cluster["names"].most_common()]
        entity = setup.add_entity(session, workspace, name=names[0])
        own_party[key] = entity.counterparty_id
        for name in names:
            setup.remember_alias(session, workspace.id, entity.counterparty_id, name, source="registry")
    # Сведённые человеком написания контрагентов — тоже псевдонимы.
    for group in decisions.get("merges") or []:
        keys = [key for key in group if key in current.clusters]
        if len(keys) < 2:
            continue
        head = current.clusters[keys[0]]["names"].most_common(1)[0][0]
        resolved = Registry(session, workspace).resolve_party(head, slot="customer")
        if resolved.party is None:
            continue
        for key in keys:
            for name in current.clusters.get(key, {"names": Tally()})["names"]:
                setup.remember_alias(session, workspace.id, resolved.party.id, name, source="registry")
    registry = Registry(session, workspace)

    # 2. Свои поля для колонок с решением «своё поле».
    custom_keys: dict[str, str] = {}
    for cid, decision in current.columns.items():
        if decision.get("action") != "custom":
            continue
        title = (decision.get("title") or "").strip() or "Колонка без названия"
        existing = next((item for item in fields_of(session, workspace.id) if norm(item.title) == norm(title)), None)
        item = existing or setup.add_field(session, workspace, title=title, type=decision.get("type") or "text")
        custom_keys[cid] = item.key
    # Шапки, сопоставленные человеком, — новые написания поля.
    _learn_names(session, workspace, current)
    registry = Registry(session, workspace)

    # 3. Статусы и виды: смыслы, назначенные на разборе.
    for spec, meaning in (decisions.get("statuses") or {}).items():
        field_key, _, value = spec.partition(":")
        if field_key in ("status", "type", "subject") and meaning:
            item = registry.resolve_value(field_key, value)
            if item is not None:
                setup.update_value(session, workspace, item.id, {"meaning": {**(item.meaning or {}), **meaning}})
    registry = Registry(session, workspace)

    # 4. Договоры: собираются в памяти (стороны, значения списков и люди
    # заводятся по дороге), вставляются пачкой.
    position = int(session.scalar(sa.select(sa.func.max(Contract.position)).where(Contract.workspace_id == workspace.id)) or 0)
    diffs = {item["id"]: item for item in next(s for s in batch.report["sections"] if s["key"] == "diffs")["items"]}
    taken_from_sheet: dict[str, dict[str, Any]] = {}
    for diff_id, item in diffs.items():
        if item["take"] == "sheet":
            taken_from_sheet.setdefault(item["main_ref"], {})[item["field"]] = item["sheet"]
    orphans_chosen = decisions.get("orphans") or {}
    end_dates = decisions.get("end_dates") or {}
    failed: list[dict[str, Any]] = []
    now = datetime.now(timezone.utc)
    rows_to_create = current.main_rows + [row for row in current.orphans if orphans_chosen.get(row["ref"], True)]
    created = 0
    chunk: list[tuple[Contract, list[Employee]]] = []

    def insert_chunk() -> None:
        # Пачкой в базу и сразу из сессии: 20 000 объектов договоров,
        # державшихся до конца заведения, поднимали процесс на сотню мегабайт.
        nonlocal created
        if not chunk:
            return
        contracts = [contract for contract, _ in chunk]
        session.add_all(contracts)
        session.flush()
        people_rows = [
            ContractPerson(contract_id=contract.id, employee_id=employee.id, position=index)
            for contract, people in chunk
            for index, employee in enumerate(people)
        ]
        session.add_all(people_rows)
        session.flush()
        for item in [*contracts, *people_rows]:
            session.expunge(item)
        created += len(chunk)
        chunk.clear()

    for row in rows_to_create:
        if row["ref"] in current.existing:
            continue  # повторная загрузка: сверка — отдельным шагом, не молча
        values = {**row["values"], **taken_from_sheet.get(row["ref"], {})}
        try:
            chunk.append(
                _contract_from_row(
                    registry, workspace, actor, batch, values, custom_keys,
                    position=position + POSITION_STEP, now=now, end_kind=end_dates.get(row["ref"]),
                )
            )
            position += POSITION_STEP
        except FinanceError as exc:
            failed.append({"ref": row["ref"], "error": str(exc)})
        if len(chunk) >= APPLY_CHUNK:
            insert_chunk()
    insert_chunk()

    # 5. Листы: главный — по главному листу, остальные — по блокам.
    _build_views(session, workspace, registry, current, decisions, batch)

    batch.status = "applied"
    batch.applied_at = now
    batch.report = {**batch.report, "result": {"created": created, "failed": failed}}
    # Разобранные строки после заведения не нужны: протокол остаётся в
    # отчёте, файл второй раз не читается. На 20 000 строк это 6 МБ в базе.
    batch.staged = {}
    bump(session, workspace.id, "schema")

    # 6. Номер изменения — последним: блокировка счётчика держится только
    # на UPDATE партии и записи истории, а не на всём заведении.
    seq = bump(session, workspace.id, "contracts")
    if created:
        _stamp_seq(session, workspace.id, batch.id, seq)
    history.write(
        session,
        workspace,
        kind="contract.import",
        entity="contract_import",
        entity_id=batch.id,
        title=f"реестр загружен из «{batch.file_name}»: {created} {_plural(created, 'договор', 'договора', 'договоров')}",
        after={"created": created, "failed": len(failed)},
        actor=actor.email,
    )
    session.flush()
    return {"created": created, "failed": failed}


def _learn_names(session: Session, workspace: Workspace, current: Plan) -> None:
    by_key = {item.key: item for item in fields_of(session, workspace.id)}
    for block in current.blocks:
        for column in block["columns"]:
            decision = current.columns.get(f"{block['id']}#{column['index']}") or {}
            if decision.get("action") != "field" or column["key"] == decision.get("field"):
                continue
            item = by_key.get(decision.get("field"))
            header = norm(column["header"])
            if item is not None and header and header not in (item.names or []):
                item.names = list(item.names or []) + [header]
    session.flush()


def _contract_from_row(
    registry: Registry,
    workspace: Workspace,
    actor: Actor,
    batch: ContractImport,
    values: dict[str, Any],
    custom_keys: dict[str, str],
    *,
    position: int,
    now: datetime,
    end_kind: str | None,
) -> tuple[Contract, list[Employee]]:
    """Договор из строки файла — в памяти, ещё не в сессии; и его ответственные.

    Стороны, значения списков и люди, которых ещё нет, заводятся по дороге
    (это отдельные записи, их ищут следующие строки). Сам договор вставляет
    `apply` — всей партией одной пачкой; все колонки заданы явно, чтобы
    пачка была одним INSERT, а не группами по набору заполненных колонок.

    Строку целиком отказывает только снимок оплат, который не читается как
    сумма; он проверяется первым, до того как строка что-то завела в базе.
    Остальное нечитаемое остаётся текстом с замечанием у договора.
    """
    snapshot: dict[str, Any] = {}
    for key in SNAPSHOT_FIELDS:
        raw = values.get(key)
        if raw in (None, ""):
            continue
        amount, terms = read_money(raw, field="Сумма") if looks_numeric(str(raw)) else (None, str(raw))
        snapshot[key.replace("_snapshot", "")] = _plain(amount) if amount is not None else terms
    contract = Contract(
        id=uuid.uuid4(),
        workspace_id=workspace.id,
        number="",
        number_key="",
        signed_at=None,
        executor_id=None,
        customer_id=None,
        type_id=None,
        subject_id=None,
        status_id=None,
        economic_role_id=None,
        department_id=None,
        billing="",
        amount=None,
        amount_terms="",
        currency="KZT",
        planned_end_at=None,
        end_date=None,
        end_kind="",
        folder_url="",
        amendments_text="",
        amendments_summary_text="",
        note="",
        attrs={},
        file_snapshot={},
        provenance={},
        acknowledged={},
        field_seq={},
        position=position,
        seq=0,
        source="import",
        import_id=batch.id,
        created_by=actor.user_id,
        updated_by=actor.user_id,
        created_at=now,
        updated_at=now,
        deleted_at=None,
    )
    people: dict[str, list[Employee]] = {}
    changed: set[str] = set()
    for key, raw in values.items():
        if key in SNAPSHOT_FIELDS:
            continue
        if key.startswith("custom:"):
            field_key = custom_keys.get(key[len("custom:"):])
            if field_key:
                contract.attrs = {**(contract.attrs or {}), field_key: str(raw)}
                changed.add(field_key)
            continue
        if key == "department" and re.search(r"[,;\n]", str(raw)):
            # «ОБО, НО, ЮО, HR» — не отдел, а подсказка из шапки: не угадываем.
            _keep_raw(contract, key, raw)
            continue
        try:
            _set_field(contract, key, raw, registry, people_out=people)
        except FinanceError:
            # Одна нечитаемая ячейка («12q» в дате) не роняет строку: договор
            # заводится, текст сохраняется, у договора — замечание.
            _keep_raw(contract, key, raw)
            continue
        changed.add(key)
    if snapshot:
        contract.file_snapshot = {**snapshot, "as_of": now.date().isoformat(), "file": batch.file_name}
    _derive(contract, registry, changed | {"type", "subject", "executor", "customer", "end_date"})
    if end_kind in ("terminated", "fulfilled", "unknown") and contract.end_date is not None:
        contract.end_kind = end_kind
        contract.provenance = {**(contract.provenance or {}), "end_kind": "manual"}
    # Номер изменения проставит `_stamp_seq` всей партии; ключи — уже сейчас.
    contract.field_seq = {key: 0 for key in sorted(changed | {"billing", "economic_role", "end_kind"})}
    return contract, list(people.get("people") or [])


def _stamp_seq(session: Session, workspace_id: uuid.UUID, batch_id: uuid.UUID, seq: int) -> None:
    """Один UPDATE: номер изменения партии — в `seq` и во все ключи `field_seq`."""
    table = Contract.__table__
    if session.get_bind().dialect.name == "postgresql":
        field_seq = sa.text(
            "(SELECT COALESCE(jsonb_object_agg(k.key, to_jsonb(CAST(:stamp AS BIGINT))), '{}'::jsonb) "
            "FROM jsonb_object_keys(field_seq) AS k(key))"
        )
    else:
        field_seq = sa.text("(SELECT COALESCE(json_group_object(key, :stamp), '{}') FROM json_each(field_seq))")
    session.execute(
        sa.update(table)
        .where(table.c.workspace_id == workspace_id, table.c.import_id == batch_id)
        .values(seq=seq, field_seq=field_seq.bindparams(stamp=seq))
    )


def _keep_raw(contract: Contract, key: str, raw: Any) -> None:
    attrs = dict(contract.attrs or {})
    kept = dict(attrs.get(RAW_KEY) or {})
    kept[key] = str(raw)
    attrs[RAW_KEY] = kept
    contract.attrs = attrs


def _build_views(
    session: Session,
    workspace: Workspace,
    registry: Registry,
    current: Plan,
    decisions: dict[str, Any],
    batch: ContractImport,
) -> None:
    rules = {item["block"]: item for item in next(s for s in batch.report["sections"] if s["key"] == "rules")["items"]}
    main_view = next((view for view in registry.views if view.main), None)
    by_sheet: dict[str, list[dict[str, Any]]] = {}
    for block in current.blocks:
        by_sheet.setdefault(block["sheet"], []).append(block)
    field_keys = set(registry.field_by_key)
    for sheet, blocks in by_sheet.items():
        view_blocks = []
        for block in blocks:
            columns = []
            for column in block["columns"]:
                decision = current.columns.get(f"{block['id']}#{column['index']}") or {}
                key = None
                if column["key"] == ROW_NUMBER:
                    key = ROW_NUMBER
                elif decision.get("action") == "field":
                    key = decision.get("field")
                elif decision.get("action") == "custom":
                    title = (decision.get("title") or "").strip() or "Колонка без названия"
                    key = next((item.key for item in registry.fields if norm(item.title) == norm(title)), None)
                if key and (key in field_keys or key == ROW_NUMBER):
                    columns.append({"key": key, "label": _label(column["header"]), "width": column.get("width")})
            if sheet == current.main_sheet:
                rule_filter: dict[str, Any] = {"any": []}
                defaults: dict[str, Any] = {}
            else:
                rule_filter, defaults = _ids_filter(rules.get(block["id"], {}).get("filter") or {"any": []}, registry)
            roles = dict(block.get("roles") or {})
            if roles.get("executor") and re.search(r"займодав", roles["executor"], re.IGNORECASE):
                financing = registry.economic.get("financing")
                if financing is not None:
                    defaults.setdefault("economic_role", str(financing.id))
            view_blocks.append(
                {
                    "title": block["title"],
                    "filter": rule_filter,
                    "roles": roles,
                    "columns": columns,
                    "defaults": defaults,
                }
            )
        style = {"header_fill": next((b["header_fill"] for b in blocks if b.get("header_fill")), ""), "source": batch.file_name}
        if sheet == current.main_sheet and main_view is not None:
            setup.upsert_view(session, workspace, {"title": sheet, "blocks": view_blocks, "style": style}, main_view.id)
        else:
            existing = next((view for view in registry.views if norm(view.title) == norm(sheet) and not view.main), None)
            if existing is not None:
                setup.upsert_view(session, workspace, {"blocks": view_blocks, "style": style}, existing.id)
            else:
                setup.upsert_view(session, workspace, {"title": sheet, "blocks": view_blocks, "style": style})


def _ids_filter(rule_filter: dict[str, Any], registry: Registry) -> tuple[dict[str, Any], dict[str, Any]]:
    """Фильтр с написаниями → фильтр с идентификаторами; и подстановки кармана.

    Подстановка ставится, только если условие однозначно: «вид ∈ {Аренда}» —
    карман блока ставит «Аренда»; «вид ∈ {Абонентское, Разовая}» — не ставит
    ничего, выбирать между ними за человека нельзя.
    """
    groups = []
    defaults: dict[str, Any] = {}
    for index, group in enumerate(rule_filter.get("any") or []):
        items = []
        for condition in group.get("all") or []:
            key = condition["field"]
            if condition.get("op") == "is":
                items.append(dict(condition))
                if index == 0 and key == "executor_is_own" and condition.get("value"):
                    defaults["own_side"] = "executor"
                if index == 0 and key == "customer_is_own" and condition.get("value"):
                    defaults.setdefault("own_side", "customer")
                continue
            ids = []
            for text in condition.get("value") or []:
                if key == "department":
                    item = registry.resolve_department(text)
                else:
                    item = registry.resolve_value(key, text)
                if item is not None:
                    ids.append(str(item.id))
            items.append({"field": key, "op": "in", "value": ids})
            if index == 0 and len(ids) == 1 and key in ("type", "subject", "department"):
                defaults[key] = ids[0]
        groups.append({"all": items})
    return {"any": groups}, defaults


def cancel(session: Session, workspace: Workspace, batch_id: uuid.UUID) -> None:
    batch = get_batch(session, workspace, batch_id)
    if batch.status == "preview":
        batch.status = "cancelled"
        batch.staged = {}
        session.flush()


__all__ = [
    "HeaderMatcher",
    "ImportFailed",
    "RULE_STRAY_ROWS",
    "RULE_STRAY_SHARE",
    "analyze",
    "apply",
    "cancel",
    "decide",
    "find_blocks",
    "get_batch",
    "match_header",
    "plan",
    "read_workbook",
    "report",
    "rule_limit",
    "start",
    "suggest_rule",
]
