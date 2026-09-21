"""Пределы значений: отказ с объяснением вместо пятисотой и вместо тихой цифры.

Каждый тест здесь — строка из стресс-прогона 21 сентября 2026, который
закончился пятисотой ошибкой, расхождением в копейку или экраном, обещавшим не
то, что случится.
"""
from __future__ import annotations

from datetime import date
from decimal import Decimal

import pytest

from app.finance import importing, invoices, recurring, service
from app.finance.db import finance_session
from app.finance.service import FinanceError


@pytest.fixture
def finance_db(tmp_path, monkeypatch):
    from sqlalchemy import create_engine

    from app.finance import db as finance_db_module

    engine = create_engine(f"sqlite:///{tmp_path / 'finance.db'}", future=True)
    monkeypatch.setattr(finance_db_module, "get_engine", lambda: engine)
    monkeypatch.setattr(finance_db_module, "_session_factory", None)
    monkeypatch.setattr(finance_db_module, "_initialized", False)
    yield engine


def _cash(session, space):
    return next(a for a in service.list_accounts(session, space.id) if a.name == "Касса")


def _income(session, space, **changes):
    data = dict(kind="income", paid_at=date(2026, 9, 3), amount=Decimal("100"), account_to_id=_cash(session, space).id)
    data.update(changes)
    return service.create_operation(session, space, service.OperationInput(**data))


def test_sorok_devyatok_otkaz_a_ne_500(finance_db):
    """`numeric(18, 2)` переполнялся в Postgres уже после проверок."""
    with finance_session() as session:
        space = service.ensure_workspace(session)
        with pytest.raises(FinanceError) as exc:
            _income(session, space, amount=Decimal("9" * 40))
    assert "слишком большая" in str(exc.value)


def test_kurs_perepolnyaet_summu_v_valyute_kompanii(finance_db):
    with finance_session() as session:
        space = service.ensure_workspace(session)
        with pytest.raises(FinanceError):
            _income(session, space, amount=Decimal("900000000000"), rate=Decimal("1000"))


def test_dolya_kopeyki_okruglyaetsya_odinakovo_vezde(finance_db):
    """0.005 весило копейку в сводке и ноль в `amount_base` — выписка расходилась."""
    with finance_session() as session:
        space = service.ensure_workspace(session)
        operation = _income(session, space, amount=Decimal("0.015"))
        assert operation.amount == Decimal("0.02")
        assert operation.amount_base == operation.amount


def test_summa_menshe_kopeyki_otkaz(finance_db):
    with finance_session() as session:
        space = service.ensure_workspace(session)
        with pytest.raises(FinanceError) as exc:
            _income(session, space, amount=Decimal("0.004"))
    assert "меньше копейки" in str(exc.value)


@pytest.mark.parametrize("year", [1900, 2999])
def test_god_s_opechatkoy_otkaz(finance_db, year):
    """Дата 2999 года молча становилась ожиданием и оседала в календаре навсегда."""
    with finance_session() as session:
        space = service.ensure_workspace(session)
        with pytest.raises(FinanceError) as exc:
            _income(session, space, paid_at=date(year, 1, 1))
    assert "опечатку" in str(exc.value)


def test_pravka_ozhidaniya_bez_scheta(finance_db):
    """Ожидание без счёта разрешено — и правка не должна проверять его как факт."""
    with finance_session() as session:
        space = service.ensure_workspace(session)
        plan = _income(session, space, status="plan", account_to_id=None, paid_at=date(2026, 12, 1))
        updated = service.update_operation(session, space, plan.id, {"comment": "уточнили"})
        assert updated.comment == "уточнили"


def test_nachalnyy_ostatok_perepolnenie(finance_db):
    with finance_session() as session:
        space = service.ensure_workspace(session)
        with pytest.raises(FinanceError):
            service.create_account(session, space, name="Огромный", starting_balance="9" * 30)


def test_schet_faktura_ogromnaya_tsena(finance_db):
    with finance_session() as session:
        space = service.ensure_workspace(session)
        with pytest.raises(FinanceError):
            invoices.create(
                session, space, kind="out", issued_at=date(2026, 9, 1), due_at=date(2026, 10, 1),
                lines=[{"title": "т", "quantity": "1", "price": "9" * 25}],
            )
        with pytest.raises(FinanceError):
            invoices.create(
                session, space, kind="out", issued_at=date(2026, 9, 1), due_at=date(2026, 10, 1),
                lines=[{"title": "т", "quantity": "много", "price": "10"}],
            )


def test_povtorenie_s_otritsatelnym_dnem(finance_db):
    """День −5 записывался приведённым, а первая дата считалась из сырого — 500."""
    with finance_session() as session:
        space = service.ensure_workspace(session)
        rule = recurring.create(
            session, space, title="аренда", kind="expense", amount=Decimal("100"),
            period="month", day=-5, start_at=date(2026, 9, 10), account_from_id=_cash(session, space).id,
        )
        assert rule.day == 1
        assert rule.next_at == date(2026, 10, 1)


def test_predprosmotr_povtora_schitaet_povtory(finance_db):
    """Экран решения обещал «готово 2» там, где заведётся ноль."""
    text = "Дата;Сумма;На счёт;Комментарий\n01.09.2026;100;Касса;а\n02.09.2026;200;Касса;б\n"
    with finance_session() as session:
        space = service.ensure_workspace(session)
        names = [a.name for a in service.list_accounts(session, space.id)]
        first = importing.analyze(text.encode("utf-8"), "f.csv", names)
        batch = service.save_preview(session, space, first)
        service.apply_batch(session, space, batch.id)

        again = importing.analyze(text.encode("utf-8"), "f.csv", names)
        service.save_preview(session, space, again)
        assert again.counts["ready"] == 0
        assert again.counts["duplicate"] == 2
        assert sum(again.counts[key] for key in ("ready", "failed", "skipped", "duplicate")) == again.counts["total"]


def test_itogi_zhurnala_po_vsemu_filtru_a_ne_po_stranitse(finance_db):
    """Карточки над журналом складывали 250 строк страницы — врали в 20 раз."""
    with finance_session() as session:
        space = service.ensure_workspace(session)
        for day in range(1, 11):
            _income(session, space, paid_at=date(2026, 9, day), amount=Decimal("100"))
        _income(session, space, paid_at=date(2026, 12, 1), amount=Decimal("999"), status="plan")
        page, total = service.list_operations(session, space.id, limit=3)
        sums = service.operation_sums(session, space.id)
    assert len(page) == 3 and total == 11
    # Все десять фактов, без ожидания: оно ещё не поступило.
    assert sums["income"] == Decimal("1000.00")


def test_vygruzka_v_excel_skhoditsya_s_itogami(finance_db):
    """Файл открывается, суммы — числа, приход минус расход — как в отчёте."""
    import io

    from openpyxl import load_workbook

    from app.finance import export

    with finance_session() as session:
        space = service.ensure_workspace(session)
        cash = _cash(session, space)
        bank = next(a for a in service.list_accounts(session, space.id) if a.id != cash.id)
        _income(session, space, amount=Decimal("1500.50"), comment="=Оплата по счёту")
        service.create_operation(session, space, service.OperationInput(
            kind="expense", paid_at=date(2026, 9, 4), amount=Decimal("500"), account_from_id=cash.id))
        service.create_operation(session, space, service.OperationInput(
            kind="transfer", paid_at=date(2026, 9, 5), amount=Decimal("300"),
            account_from_id=cash.id, account_to_id=bank.id))
        data, count = export.journal_xlsx(session, space, service.OperationFilter())
    assert count == 3
    sheet = load_workbook(io.BytesIO(data)).active
    rows = list(sheet.iter_rows(min_row=2, max_row=4, values_only=True))
    income = sum(row[2] or 0 for row in rows)
    expense = sum(row[3] or 0 for row in rows)
    transfer = sum(row[4] or 0 for row in rows)
    assert (income, expense, transfer) == (1500.5, 500, 300)
    # Назначение платежа с «=» осталось текстом, а не стало формулой.
    assert any(row[11] == "=Оплата по счёту" for row in rows)
    assert sheet.cell(row=6, column=3).value.startswith("=SUBTOTAL")


def test_nachalnyy_ostatok_menyaetsya_i_otmenyaetsya(finance_db):
    """Остаток ставился только при создании счёта — у заведённых сами не менялся."""
    from app.finance import history

    with finance_session() as session:
        space = service.ensure_workspace(session)
        cash = _cash(session, space)
        account, before = service.set_starting_balance(session, space, cash.id, "28012.68")
        entry = history.write(
            session, space, kind="account.balance", entity="account", entity_id=account.id,
            before={"starting_balance": str(before)}, after={"starting_balance": str(account.starting_balance)},
        )
        assert account.starting_balance == Decimal("28012.68")
        history.undo(session, space, entry.id)
        assert session.get(type(account), account.id).starting_balance == Decimal("0")


def test_sverka_vypiski_s_bankom(finance_db):
    """Остатки банка против разобранных строк; начальный остаток — из выписки."""
    from app.finance.importing import DateReading, ParsedRow, Preview

    rows = [
        ParsedRow(line=2, raw={}, state="imported", values={"amount": "1000.00", "kind": "income"}),
        ParsedRow(line=3, raw={}, state="imported", values={"amount": "250.50", "kind": "expense"}),
    ]
    with finance_session() as session:
        space = service.ensure_workspace(session)
        preview = Preview(
            file_name="g.pdf", header_line=0, mapping={}, unused_columns=[], rows=rows,
            date_reading=DateReading("dmy"),
            bank={"period_start": "2026-09-01", "period_end": "2026-09-30",
                  "opening_balance": "500.00", "closing_balance": "1249.50", "account": "Касса"},
        )
        out = service.reconcile_statement(session, space, preview)
        assert out["file_net"] == "749.50"
        assert out["gap"] == "0.00"
        assert out["can_set_start"] is True
        # Раньше периода по счёту была операция — остаток на начало задан ею,
        # и начальный остаток из выписки уже не предлагается.
        _income(session, space, paid_at=date(2026, 8, 1), amount=Decimal("500"))
        again = service.reconcile_statement(session, space, preview)
        assert again["earlier_operations"] == 1
        assert again["ledger_opening"] == "500.00"
        assert again["can_set_start"] is False


def test_kaspi_gold_ostatki_po_date_a_ne_po_poryadku():
    """Первое «Доступно на» в шапке — остаток на КОНЕЦ, а бралось за начало."""
    from app.services.document_service import _extract_pdf_metadata

    lines = [
        "ВЫПИСКА", "по Kaspi Gold за период с 18.09.25 по 18.09.26",
        "Доступно на 18.09.26:", "+ 21 439,09 ₸",
        "Краткое содержание операций по карте:",
        "Доступно на 18.09.25", "+ 28 012,68 ₸",
        "Доступно на 18.09.26", "+ 21 439,09 ₸",
    ]
    meta = _extract_pdf_metadata("g.pdf", lines, [])
    assert meta.opening_balance == 28012.68
    assert meta.closing_balance == 21439.09


def test_naznachenie_eto_kommentariy_a_ne_statya():
    """«Назначение» из банковской выгрузки заводило по статье на каждую строку."""
    text = "Дата;Сумма;Назначение;На счёт\n01.09.2026;125000;Оплата по счёту 311;Касса\n"
    preview = importing.analyze(text.encode("utf-8"), "bank.csv", ["Касса"], date_order="dmy")
    values = preview.rows[0].values
    assert values["comment"] == "Оплата по счёту 311"
    assert not values.get("category")
