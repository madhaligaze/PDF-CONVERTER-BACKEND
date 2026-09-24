"""Загрузка реестра: правила листов, совпадения по номеру, партия одним номером.

Что здесь держится и почему
───────────────────────────
* Отбор без условий у неглавного листа не отбирает ничего. Раньше пустое
  правило значило «все договоры», и блок, которому правило не подобралось,
  показывал весь реестр (лист на 192 строки — пять тысяч договоров).
* Предложенное правило, которое тянет в лист много лишнего, без человека не
  принимается: раздел «Правила листов» блокирует «Завести», пока по блоку нет
  решения — принять как есть, своё правило или пустой блок.
* Строка листа с тем же номером, что у чужого договора, но без единой общей
  стороны, — отдельный договор, а не «совпало по номеру»: иначе она
  приклеивалась к чужому договору и не заводилась вовсе.
* Партия получает один номер изменения, взятый в конце, а разобранные строки
  после заведения не хранятся.
* Поиск поля по шапке отвечает так же, как прежний перебор.

Файлы собираются здесь же, по форме реестра ЮО BBC, без его данных.
"""
from __future__ import annotations

import io
import random
import re
from datetime import datetime
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook

from app.books.layout import norm, squash
from app.finance import service as finance_service
from app.finance.contracts import export, importer, service, views
from app.finance.contracts.fields import SYSTEM_FIELDS, current
from app.finance.contracts.models import Contract, ContractImport
from app.finance.db import finance_session
from app.finance.service import FinanceError

FULL = service.Access(view=True, edit=True, setup=True)
ACTOR = service.Actor(None, "owner@test")

HEAD = [
    "№", "Текущее состояние (действующий/\nнедействующий)", "Ссылка на Битрикс 24",
    "Планируемый срок завершения", "Ответственное лицо (Ф.И.О)", "Исполнитель \nBBC",
    "Заказчик \n(клиент)", "№ Договора", "Дата заключения Договора", "Отдел\n(ОБО, НО,\n ЮО)",
    "Вид услуги", "Предмет Договора (взыскание долга, и т.д.)", "Сумма Договора",
    "Оплачено на текущую дату", "Остаток оплаты", "№ Дополнительное соглашение/дата ",
    "Предмет Доп.соглашения", "Дата расторжения Договора/дата исполения разового Договора", "*Примечания",
]


def _row(n, executor, customer, number, kind, subject, amount=100000, status="действующий", signed=None):
    return [n, status, None, None, None, executor, customer, number, signed or datetime(2025, 1, n % 28 + 1),
            "ЮО", kind, subject, amount, None, None, None, None, None, None]


def _head(executor_title: str, customer_title: str) -> list[str]:
    head = list(HEAD)
    head[5], head[6] = executor_title, customer_title
    return head


@pytest.fixture
def finance_db(tmp_path, monkeypatch):
    from sqlalchemy import create_engine

    from app.finance import db as finance_db_module

    engine = create_engine(f"sqlite:///{tmp_path / 'finance.db'}", future=True)
    monkeypatch.setattr(finance_db_module, "get_engine", lambda: engine)
    monkeypatch.setattr(finance_db_module, "_session_factory", None)
    monkeypatch.setattr(finance_db_module, "_initialized", False)
    yield engine


def _section(batch, key):
    return next(section for section in batch.report["sections"] if section["key"] == key)


def _book(main_rows: list[list], blocks: list[tuple[str, str, str, str, list[list]]]) -> bytes:
    """«Сводная» и блоки других листов: (лист, название, исполнитель, заказчик, строки)."""
    book = Workbook()
    main = book.active
    main.title = "Сводная"
    main.append(HEAD)
    for row in main_rows:
        main.append(row)
    sheets: dict[str, object] = {}
    for sheet_name, title, executor_title, customer_title, rows in blocks:
        sheet = sheets.get(sheet_name)
        if sheet is None:
            sheet = sheets[sheet_name] = book.create_sheet(sheet_name)
        else:
            sheet.append([])
        sheet.append([None, None, None, title])
        sheet.append(_head(executor_title, customer_title))
        for row in rows:
            sheet.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _entities_and_columns(batch) -> dict:
    columns = {
        column["id"]: {"action": "skip"}
        for block in _section(batch, "columns")["items"]
        for column in block["columns"]
        if column["decision"].get("action") == "ask"
    }
    own = [item["key"] for item in _section(batch, "entities")["items"] if item["own"]]
    return {"columns": columns, "entities": {"confirmed": True, "own": own}}


def _wide_book() -> bytes:
    """Блок, которому правило подбирается только «вид ∈ {Иное}»: лишних много.

    В блоке «Заказчик ГК» наше юрлицо стоит то заказчиком, то исполнителем
    (так бывает в живом реестре), предметов шесть — больше, чем правило
    перечисляет, — и признаков, общих для строк блока и редких в реестре, нет,
    кроме вида «Иное». А «Иное» в главном листе у двенадцати агентских.
    """
    main = [_row(n, "BBC", f"ТОО Клиент {n}", f"№ АГ-{n}", "Иное", "Агентский") for n in range(1, 13)]
    main += [_row(20 + n, "BBC", f"ТОО Абонент {n}", f"№ АБ-{n}", "Абонентское обслуживание",
                  "Бухгалтерское сопровождение") for n in range(1, 6)]
    block = [
        _row(n, "BBC" if n % 2 else f"ТОО Поставщик {n}", f"ТОО Поставщик {n}" if n % 2 else "BBC",
             f"№ ЗК-{n}", "Иное", subject)
        for n, subject in enumerate(WIDE_SUBJECTS, start=1)
    ]
    return _book(main, [("Заказчик ГК", "Заказчик ГК", "Заказчик", "Исполнитель", block)])


WIDE_SUBJECTS = ["Маркетинг", "Реклама", "Оценка", "Перевод", "Печать", "Уборка"]


# ── Пустое правило ────────────────────────────────────────────────────────────


def test_pustoe_pravilo_ne_otbiraet_nichego_krome_glavnogo_lista():
    facts = {"type": "t1"}
    empty = {"any": []}
    rule = {"any": [{"all": [{"field": "type", "op": "in", "value": ["t1"]}]}]}
    main = SimpleNamespace(key="main", main=True, blocks=[{"filter": empty}])
    main_two = SimpleNamespace(key="main2", main=True, blocks=[{"filter": rule}, {"filter": empty}])
    main_ruled = SimpleNamespace(key="main3", main=True, blocks=[{"filter": {"any": [{"all": [
        {"field": "type", "op": "in", "value": ["other"]}]}]}}])
    other = SimpleNamespace(key="other", main=False, blocks=[{"filter": empty}])
    other_two = SimpleNamespace(key="other2", main=False, blocks=[{"filter": empty}, {"filter": rule}])

    assert views.matches(empty, facts) is False
    assert views.place(main, facts) == 0
    assert views.place(main_two, facts) == 0
    assert views.place(main_two, {"type": "t2"}) == 1
    # Главный лист держит договор, даже если ни одно правило его блоков не подошло.
    assert views.place(main_ruled, facts) == 0
    assert views.place(other, facts) is None
    assert views.place(other_two, facts) == 1
    assert views.membership(facts, [main, other, other_two]) == [
        {"view": "main", "block": 0}, {"view": "other2", "block": 1}
    ]


# ── Правило с лишними ждёт человека ───────────────────────────────────────────


def test_netochnoe_pravilo_blokiruet_zavedenie(finance_db):
    with finance_session() as session:
        workspace = finance_service.ensure_workspace(session)
        batch = importer.start(session, workspace, ACTOR, _wide_book(), "реестр.xlsx")
        batch = importer.decide(session, workspace, batch.id, _entities_and_columns(batch))
        rules = _section(batch, "rules")
        item = next(item for item in rules["items"] if item["block"] == "Заказчик ГК#0")
        assert item["in_sheet"] == 6
        assert item["filter"]["any"][0]["all"][0]["field"] == "type"
        assert item["extra"] == 12 > importer.rule_limit(6)
        assert item["needs_decision"] and item["reason"]
        assert rules["pending"] == ["Заказчик ГК#0"]
        assert "rules" in batch.report["blocking"]
        with pytest.raises(FinanceError, match="Правила листов"):
            importer.apply(session, workspace, FULL, ACTOR, batch.id)


def test_prinyat_kak_est_zapisyvaet_pravilo_v_reshenie(finance_db):
    with finance_session() as session:
        workspace = finance_service.ensure_workspace(session)
        batch = importer.start(session, workspace, ACTOR, _wide_book(), "реестр.xlsx")
        decisions = _entities_and_columns(batch)
        decisions["rules"] = {"Заказчик ГК#0": {"action": "accept"}}
        batch = importer.decide(session, workspace, batch.id, decisions)
        suggested = next(item for item in _section(batch, "rules")["items"] if item["block"] == "Заказчик ГК#0")
        # Принятое правило записано в решение — следующее решение его не подменит.
        frozen = batch.decisions["rules"]["Заказчик ГК#0"]
        assert frozen["action"] == "accept" and frozen["filter"] == suggested["filter"]
        assert suggested["source"] == "accepted" and not suggested["needs_decision"]
        assert "rules" not in batch.report["blocking"]
        result = importer.apply(session, workspace, FULL, ACTOR, batch.id)
        assert result["created"] == 23 and not result["failed"]
        listing = service.list_all(session, workspace, FULL)
        view = next(v for v in service.Registry(session, workspace).views if v.title == "Заказчик ГК")
        members = [item for item in listing["contracts"] if any(m["view"] == view.key for m in item["views"])]
        # Принято как есть — с лишними: 6 строк листа и 12 агентских.
        assert suggested["extra"] == 12 and len(members) == 18


def test_pustoy_blok_i_svoe_pravilo(finance_db):
    with finance_session() as session:
        workspace = finance_service.ensure_workspace(session)
        batch = importer.start(session, workspace, ACTOR, _wide_book(), "реестр.xlsx")
        decisions = _entities_and_columns(batch)
        bad = {"any": [{"all": [{"field": "number", "op": "in", "value": ["№ ЗК-1"]}]}]}
        with pytest.raises(FinanceError, match="отбирать нельзя"):
            importer.decide(session, workspace, batch.id, {**decisions, "rules": {"Заказчик ГК#0": {"action": "rule", "filter": bad}}})
        own_rule = {"any": [{"all": [{"field": "subject", "op": "in", "value": WIDE_SUBJECTS}]}]}
        batch = importer.decide(session, workspace, batch.id, {**decisions, "rules": {"Заказчик ГК#0": {"action": "rule", "filter": own_rule}}})
        item = next(item for item in _section(batch, "rules")["items"] if item["block"] == "Заказчик ГК#0")
        assert (item["source"], item["caught"], item["extra"]) == ("manual", 6, 0)

        batch = importer.decide(session, workspace, batch.id, {"rules": {"Заказчик ГК#0": {"action": "empty"}}})
        item = next(item for item in _section(batch, "rules")["items"] if item["block"] == "Заказчик ГК#0")
        assert item["source"] == "empty" and item["filter"] == {"any": []}
        assert "rules" not in batch.report["blocking"]
        importer.apply(session, workspace, FULL, ACTOR, batch.id)
        registry = service.Registry(session, workspace)
        view = next(v for v in registry.views if v.title == "Заказчик ГК")
        assert view.blocks[0]["filter"] == {"any": []}
        listing = service.list_all(session, workspace, FULL)
        # Пустой блок — пустой лист, а не весь реестр.
        assert not [item for item in listing["contracts"] if any(m["view"] == view.key for m in item["views"])]
        assert all(any(m["view"] == "main" for m in item["views"]) for item in listing["contracts"])
        data = export.build(session, workspace, FULL, ACTOR)
    sheet = load_workbook(io.BytesIO(data))["Заказчик ГК"]
    # Название блока и шапка, строк договоров нет.
    assert sheet.max_row == 2


def test_tochnoe_pravilo_ne_sprashivaet(finance_db):
    main = [_row(n, "BBCA", f"ТОО Арендатор {n}", f"№ R-{n}", "Аренда", "Аренда нежилого помещения") for n in range(1, 9)]
    main += [_row(10 + n, "BBC", f"ТОО Клиент {n}", f"№ C-{n}", "Разовая услуга", "Аудит") for n in range(1, 9)]
    rent = [_row(n, "BBCA", f"ТОО Арендатор {n}", f"№ R-{n}", "Аренда", "Аренда нежилого помещения") for n in range(1, 9)]
    data = _book(main, [("Прочие договоры", "АРЕНДА", "Арендадатель\nBBC", "Арендатор (клиент)", rent)])
    with finance_session() as session:
        workspace = finance_service.ensure_workspace(session)
        batch = importer.start(session, workspace, ACTOR, data, "реестр.xlsx")
        batch = importer.decide(session, workspace, batch.id, _entities_and_columns(batch))
        item = _section(batch, "rules")["items"][0]
        assert (item["caught"], item["extra"], item["needs_decision"]) == (8, 0, False)
        assert batch.report["blocking"] == []


def test_porog_pravila():
    # Пять строк или десятая часть листа — что меньше.
    assert [importer.rule_limit(n) for n in (0, 1, 9, 10, 21, 49, 50, 341, 20000)] == [0, 0, 0, 1, 2, 4, 5, 5, 5]


# ── Совпадения по номеру ──────────────────────────────────────────────────────


def test_nomer_bez_obshchey_storony_ne_skleivaetsya(finance_db):
    main = [
        _row(1, "BBC", "ТОО Альфа", "№ 7", "Разовая услуга", "Аудит"),
        _row(2, "BBC", "ТОО Бета", "№ 8", "Разовая услуга", "Аудит"),
        _row(3, "BBC", "ТОО Гамма", "№ 9", "Разовая услуга", "Аудит"),
        _row(4, "BBC", "ТОО Дельта", "№ 10", "Разовая услуга", "Аудит"),
    ]
    sheet = [
        # Тот же номер, стороны переставлены — тот же договор.
        _row(1, "ТОО Альфа", "BBC", "№ 7", "Разовая услуга", "Аудит"),
        # Тот же номер и заказчик, исполнитель другой — тот же договор (сменили ТОО).
        _row(2, "BBCA", "ТОО Бета", "№ 8", "Разовая услуга", "Аудит"),
        # Тот же номер, ни одной общей стороны — другой договор.
        _row(3, "ТОО Продавец", "ИП Покупатель", "№ 9", "Разовая услуга", "Аудит"),
        # Тот же номер, общая только наша сторона — она почти в каждом договоре,
        # это не довод: другой договор.
        _row(4, "BBC", "ТОО Эпсилон", "№ 10", "Разовая услуга", "Аудит"),
    ]
    data = _book(main, [("Исполнитель ГК", "Исполнитель ГК", "Исполнитель \nBBC", "Заказчик \n(клиент)", sheet)])
    with finance_session() as session:
        workspace = finance_service.ensure_workspace(session)
        batch = importer.start(session, workspace, ACTOR, data, "реестр.xlsx")
        decisions = _entities_and_columns(batch)
        decisions["rules"] = {"Исполнитель ГК#0": {"action": "accept"}}
        batch = importer.decide(session, workspace, batch.id, decisions)
        orphans = _section(batch, "orphans")
        kinds = {item["number"]: item["kind"] for item in orphans["loose"]}
        assert kinds == {"№ 7": "swapped", "№ 8": "number_party"}
        assert [item["number"] for item in orphans["number_only"]] == ["№ 9", "№ 10"]
        assert [item["number"] for item in orphans["items"]] == ["№ 9", "№ 10"]
        assert batch.report["totals"]["create"] == 6

        # «Это он» — человек сводит строку с договором главного листа.
        ref = orphans["number_only"][0]["ref"]
        batch = importer.decide(session, workspace, batch.id, {"loose": {ref: "same"}})
        orphans = _section(batch, "orphans")
        assert [item["number"] for item in orphans["items"]] == ["№ 10"]
        assert {item["number"]: item["kind"] for item in orphans["loose"]}["№ 9"] == "number"
        assert batch.report["totals"]["create"] == 5


# ── Партия одним номером ──────────────────────────────────────────────────────


def test_partiya_odnim_nomerom_i_bez_razobrannyh_strok(finance_db):
    with finance_session() as session:
        workspace = finance_service.ensure_workspace(session)
        before = current(session, workspace.id, "contracts")
        batch = importer.start(session, workspace, ACTOR, _wide_book(), "реестр.xlsx")
        decisions = _entities_and_columns(batch)
        decisions["rules"] = {"Заказчик ГК#0": {"action": "empty"}}
        importer.decide(session, workspace, batch.id, decisions)
        result = importer.apply(session, workspace, FULL, ACTOR, batch.id)
        batch_id, space_id = batch.id, workspace.id
    with finance_session() as session:
        seq = current(session, space_id, "contracts")
        assert seq == before + 1
        contracts = session.query(Contract).filter_by(workspace_id=space_id).all()
        assert len(contracts) == result["created"] == 23
        assert {item.seq for item in contracts} == {seq}
        for item in contracts:
            assert item.field_seq and set(item.field_seq.values()) == {seq}
            assert {"billing", "economic_role", "end_kind", "number", "executor"} <= set(item.field_seq)
        stored = session.get(ContractImport, batch_id)
        assert stored.staged == {} and stored.status == "applied"
        assert stored.report["result"]["created"] == 23
        # Опрос «после номера до загрузки» видит всю партию.
        workspace = finance_service.get_workspace(session, space_id)
        assert len(service.changes(session, workspace, FULL, before)["contracts"]) == 23


def test_otkaz_stroki_ne_zavodit_ee_storony(finance_db):
    main = [_row(1, "BBC", "ТОО Альфа", "№ 1", "Разовая услуга", "Аудит"),
            _row(2, "BBC", "ТОО Бета", "№ 2", "Разовая услуга", "Аудит")]
    bad = _row(3, "BBC", "ТОО Одноразовый", "№ 3", "Разовая услуга", "Аудит")
    bad[13] = "-500"  # оплачено: отрицательная сумма — строка не заводится
    main.append(bad)
    with finance_session() as session:
        workspace = finance_service.ensure_workspace(session)
        batch = importer.start(session, workspace, ACTOR, _book(main, []), "реестр.xlsx")
        importer.decide(session, workspace, batch.id, _entities_and_columns(batch))
        result = importer.apply(session, workspace, FULL, ACTOR, batch.id)
        assert result["created"] == 2 and [item["ref"] for item in result["failed"]] == ["Сводная!4"]
        names = {party.name for party in service.Registry(session, workspace).parties.values()}
        assert "ТОО Одноразовый" not in names


# ── Поиск поля по шапке ──────────────────────────────────────────────────────


_PAREN = re.compile(r"\([^)]*\)", re.S)


def _reference_match(text, catalog):
    """Прежний перебор — эталон для быстрого поиска."""
    full = norm(text)
    if not full:
        return None, "empty", []
    if full in ("№", "#", "n", "no", "№ п/п", "п/п"):
        return importer.ROW_NUMBER, "exact", []
    core = norm(_PAREN.sub(" ", str(text or "")))
    tiers = (
        ("exact", lambda item: full in item.names),
        ("core", lambda item: bool(core) and core in item.names),
        ("squashed", lambda item: bool(squash(core)) and squash(core) in item.squashed),
        ("loose", lambda item: importer._loose(squash(core), item.squashed)),
    )
    for how, test in tiers:
        found = [item.key for item in catalog if test(item)]
        if len(found) == 1:
            return found[0], how, []
        if len(found) > 1:
            return None, "ambiguous", found
    return None, "none", []


def test_poisk_shapki_kak_perebor():
    extra = [
        SimpleNamespace(key="istochnik", title="Источник клиента", names=["источник клиента", "источник"]),
        # Своё поле, чьё написание — начало системного: мягкий ярус двусмысленен.
        SimpleNamespace(key="dogovor_x", title="Договор", names=["договор", "дата договора (скан)"]),
        SimpleNamespace(key="n2", title="Сумма", names=["сумма"]),
    ]
    catalog = importer.field_names(list(SYSTEM_FIELDS) + extra)
    matcher = importer.HeaderMatcher(catalog)
    texts = list(HEAD) + [
        "", " ", "№", "№ п/п", "N", "Исполнитель BBC", "ИСПОЛНИТЕЛЬ  bbc", "Отдел (ОБО)", "(Отдел)",
        "Сумма", "Сумма договора (тг)", "сумма дог", "Ёмкость", "Дата", "Дата заключения", "Дата расторжения",
        "Дата исполнения разового", "Предмет", "Предмет доп", "Источник", "Источники клиентов", "Договор",
        "Договор №", "Статус договора", "Ответ лицо", "Примечания*", "***", "123", "2025-01-01",
        "Абонентское обслуживание", "ТОО «Альфа»", "BBC", "https://bitrix/1", "Займодавец", "Арендатор",
    ]
    rng = random.Random(3)
    words = [name for item in catalog for name in item.names] + ["зака", "испол", "сумм", "предм", "дата з"]
    for _ in range(400):
        base = rng.choice(words)
        texts.append(rng.choice([base.upper(), base[: rng.randint(2, len(base))], f"{base} ({rng.choice(words)})",
                                 base.replace(" ", "\n"), f" {base}. "]))
    for text in texts:
        assert matcher.match(text) == _reference_match(text, catalog), text
        assert importer.match_header(text, catalog) == _reference_match(text, catalog), text


# ── Выгрузка потоком ─────────────────────────────────────────────────────────


def test_vygruzka_potokom_derzhit_oformlenie(finance_db):
    main = [_row(n, "BBCA", f"ТОО Арендатор {n}", f"№ R-{n}", "Аренда", "Аренда нежилого помещения") for n in range(1, 4)]
    main += [_row(10 + n, "BBC", f"ТОО Клиент {n}", f"№ C-{n}", "Разовая услуга", "Аудит") for n in range(1, 4)]
    rent = [_row(n, "BBCA", f"ТОО Арендатор {n}", f"№ R-{n}", "Аренда", "Аренда нежилого помещения") for n in range(1, 4)]
    audit = [_row(10 + n, "BBC", f"ТОО Клиент {n}", f"№ C-{n}", "Разовая услуга", "Аудит") for n in range(1, 4)]
    data = _book(main, [
        ("Прочие договоры", "АРЕНДА", "Арендадатель\nBBC", "Арендатор (клиент)", rent),
        ("Прочие договоры", "АУДИТ", "Исполнитель \nBBC", "Заказчик \n(клиент)", audit),
    ])
    with finance_session() as session:
        workspace = finance_service.ensure_workspace(session)
        batch = importer.start(session, workspace, ACTOR, data, "реестр.xlsx")
        importer.decide(session, workspace, batch.id, _entities_and_columns(batch))
        importer.apply(session, workspace, FULL, ACTOR, batch.id)
        registry = service.Registry(session, workspace)
        main_view = next(v for v in registry.views if v.main)
        main_view.style = {**(main_view.style or {}), "header_fill": "D9E1F2"}
        session.flush()
        out = export.build(session, workspace, FULL, ACTOR)
    book = load_workbook(io.BytesIO(out))
    sheet = book["Сводная"]
    assert sheet.freeze_panes == "A2"
    header = [cell for cell in sheet[1]]
    assert header[0].value == "№" and all(cell.font.b for cell in header if cell.value)
    assert header[1].fill.start_color.rgb.endswith("D9E1F2")
    signed = next(i for i, cell in enumerate(header) if cell.value == "Дата заключения Договора")
    assert sheet.cell(row=2, column=signed + 1).number_format == "DD.MM.YYYY"
    assert sheet.max_row == 7  # шапка и шесть договоров
    other = book["Прочие договоры"]
    rows = [[cell.value for cell in row] for row in other.iter_rows()]
    titles = [i for i, row in enumerate(rows) if row[1] in ("АРЕНДА", "АУДИТ")]
    assert [rows[i][1] for i in titles] == ["АРЕНДА", "АУДИТ"]
    # Блок: название, шапка, три договора; между блоками — две пустые строки.
    assert titles[1] - titles[0] == 1 + 1 + 3 + 2
    assert other.freeze_panes is None


def test_vygruzka_rezhet_stroki_i_kolonki_po_pravam(finance_db):
    """Потоковая книга выгружает то же, что видит человек: скрытое поле — без
    колонки, договоры чужих юрлиц — без строк."""
    main = [_row(n, "BBC", f"ТОО Клиент {n}", f"№ B-{n}", "Разовая услуга", "Аудит") for n in range(1, 4)]
    main += [_row(10 + n, "BBCA", f"ТОО Арендатор {n}", f"№ R-{n}", "Аренда", "Аренда нежилого помещения") for n in range(1, 3)]
    with finance_session() as session:
        workspace = finance_service.ensure_workspace(session)
        batch = importer.start(session, workspace, ACTOR, _book(main, []), "реестр.xlsx")
        importer.decide(session, workspace, batch.id, _entities_and_columns(batch))
        importer.apply(session, workspace, FULL, ACTOR, batch.id)
        registry = service.Registry(session, workspace)
        bbca = next(pid for pid, party in registry.parties.items() if party.name == "BBCA")
        narrow = service.Access(view=True, entity_ids=frozenset({bbca}), hidden=frozenset({"amount"}))
        data = export.build(session, workspace, narrow, ACTOR)
    sheet = load_workbook(io.BytesIO(data))["Сводная"]
    rows = [[cell.value for cell in row] for row in sheet.iter_rows()]
    assert "Сумма Договора" not in rows[0] and "№ Договора" in rows[0]
    number = rows[0].index("№ Договора")
    assert sorted(row[number] for row in rows[1:]) == ["№ R-1", "№ R-2"]
