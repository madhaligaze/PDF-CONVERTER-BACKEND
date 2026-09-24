"""Загрузка реестра из Excel: блоки со своими шапками, ничего спорного молча.

Файл собирается здесь же, по форме реестра ЮО BBC, но без его данных:
настоящий реестр с клиентами в репозиторий не кладётся.
"""
from __future__ import annotations

import io
from datetime import datetime

import pytest
from openpyxl import Workbook

from app.finance import service as finance_service
from app.finance.contracts import export, importer, service
from app.finance.db import finance_session

FULL = service.Access(view=True, edit=True, setup=True)
ACTOR = service.Actor(None, "owner@test")

HEAD = [
    "№", "Текущее состояние (действующий/\nнедействующий)", "Ссылка на Битрикс 24",
    "Планируемый срок завершения", "Ответственное лицо (Ф.И.О)", "Исполнитель \nBBC",
    "Заказчик \n(клиент)", "№ Договора", "Дата заключения Договора", "Отдел\n(ОБО, НО,\n ЮО)",
    "Вид услуги", "Предмет Договора (взыскание долга, и т.д.)", "Сумма Договора",
    "Оплачено на текущую дату", "Остаток оплаты", "№ Дополнительное соглашение/дата ",
    "Предмет Доп.соглашения", "Дата расторжения Договора/дата исполения разового Договора", "*Примечания",
]


def _row(n, status, executor, customer, number, signed, dept, kind, subject, amount, **extra):
    row = [n, status, "https://bitrix/folder/" + str(n), extra.get("planned"), extra.get("people"),
           executor, customer, number, signed, dept, kind, subject, amount,
           extra.get("paid"), extra.get("rest"), extra.get("amend"), extra.get("amend_subject"),
           extra.get("end"), extra.get("note")]
    if "t" in extra:
        row.append(extra["t"])
    return row


def _registry_file() -> bytes:
    book = Workbook()
    main = book.active
    main.title = "Сводная"
    main.append(HEAD)
    d = datetime
    main.append(_row(1, "действующий", "BBC ", "ТОО Альфа", " № 17-BUH", d(2021, 12, 1), "ОБО",
                     "Абонентское обслуживание", "Бухгалтерское сопровождение", 250000,
                     people="Наталья Петровна", amend="№ 1 от 23.07.2024\n№2 10.12.2025",
                     amend_subject="1) изменение п. 1.2\n2) увелечение цены на 250 000", paid=250000))
    main.append(_row(2, "недействующий", "SAKOMPA-M", "ТОО Бета", "№ 60-CONSULT", d(2024, 2, 1), "НО",
                     "Абонентское обслуживание", "Налоговое сопровождение", "1\xa0200\xa0000",
                     end=d(2025, 5, 12), people="Елжас, Тимур"))
    main.append(_row(3, "исполнен", "Sakompa-M", "ТОО Гамма", "№ЮО/59", d(2024, 11, 11), "ЮО",
                     "Разовая услуга", "Взыскание долга", 1000000, end=d(2025, 4, 1), planned="до 15.07.2026"))
    main.append(_row(4, "действующий", "BBC", "ИП Дельта", "№ЮО/59", d(2025, 1, 10), "ЮО",
                     "Разовая услуга", "Юридическое сопровождение", 500000, t="у Бисултана 150 000"))
    main.append(_row(5, "нужно закрыть по бух", "BBC Marketing", "BBC", "№ AG-4", d(2025, 9, 18), None,
                     "Иное", "Агентский", "20% - по разовым,\n40% - по абон."))
    main.append(_row(6, "действующий", "BBCA", "ТОО Эпсилон", "№ 50-RENT", d(2024, 7, 12), None,
                     "Аренда", "Аренда нежилого помещения", "50 000"))
    main.append(_row(7, "действующий", "BBCA", "ТОО Зета", "№ 51-RENT", d(2024, 8, 1), None,
                     "Аренда", "Аренда нежилого помещения", 60000, end="12q"))
    main.append(_row(8, "действующий", "BBC Marketing", "ТОО Альфа", "№ 12", d(2025, 3, 1), "ОБО",
                     "Абонентское обслуживание", "Бухгалтерское сопровождение", 100000))
    main.append(_row(9, "действующий", "BBC Marketing", "ТОО Бета", "№ 13", d(2025, 3, 1), "НО",
                     "Абонентское обслуживание", "Налоговое сопровождение", 200000))

    buyer = book.create_sheet("Заказчик ГК")
    buyer.append([None, None, None, "КУПЛЯ-ПРОДАЖА"])
    head = list(HEAD)
    head[5], head[6] = "Продавец", "Покупатель"
    buyer.append(head)
    buyer.append(_row(1, "Исполнен", "Халык Актив", "BBCA", "Договор купли-продажи", d(2025, 12, 26), None,
                      "Иное", "Договор купли-продажи", 2765000))
    buyer.append([])
    buyer.append([None, None, None, "Заказчик ГК"])
    head = list(HEAD)
    head[5], head[6] = "Заказчик", "Исполнитель"
    buyer.append(head)
    # Колонки наоборот: F — заказчик (наш), G — исполнитель (внешний).
    buyer.append(_row(1, "исполнен", "BBC", "ИП «DSGroup»", "№041125", d(2025, 11, 4), None,
                      "Иное", "Маркетинг", 300000))

    other = book.create_sheet("Прочие договоры")
    other.append([None, None, None, None, "АРЕНДА"])
    head = list(HEAD)
    head[5], head[6] = "Арендадатель\nBBC", "Арендатор (клиент)"
    other.append(head)
    other.append(_row(1, "действующий", "BBCA", "ТОО Эпсилон", "№ 50-RENT", d(2024, 7, 12), None,
                      "Аренда", "Аренда нежилого помещения", 55000))
    other.append(_row(2, "действующий", "BBCA", "ТОО Зета", "№ 51-RENT", d(2024, 8, 1), None,
                      "Аренда", "Аренда нежилого помещения", 60000))
    other.append(_row(3, "действующий", "BBCA", "ТОО Эта", "№ 52-RENT", d(2024, 9, 1), None,
                      "Аренда", "Аренда нежилого помещения", 70000))
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def finance_db(tmp_path, monkeypatch):
    from sqlalchemy import create_engine

    from app.finance import db as finance_db_module

    engine = create_engine(f"sqlite:///{tmp_path / 'finance.db'}", future=True)
    monkeypatch.setattr(finance_db_module, "get_engine", lambda: engine)
    monkeypatch.setattr(finance_db_module, "_session_factory", None)
    monkeypatch.setattr(finance_db_module, "_initialized", False)
    yield engine


def _section(batch, key):
    return next(section for section in batch.report["sections"] if section["key"] == key)


def _start(session):
    workspace = finance_service.ensure_workspace(session)
    return workspace, importer.start(session, workspace, ACTOR, _registry_file(), "реестр.xlsx")


def test_bloki_i_storony_po_shapke_bloka(finance_db):
    with finance_session() as session:
        _workspace, batch = _start(session)
        blocks = {item["id"]: item for item in _section(batch, "blocks")["items"]}
    assert set(blocks) == {"Сводная#0", "Заказчик ГК#0", "Заказчик ГК#1", "Прочие договоры#0"}
    assert blocks["Заказчик ГК#0"]["title"] == "КУПЛЯ-ПРОДАЖА"
    assert blocks["Заказчик ГК#0"]["roles"]["executor"] == "Продавец"
    # Во втором блоке стороны наоборот, и это видно из порядка.
    assert blocks["Заказчик ГК#1"]["roles"]["order"] == ["customer", "executor"]
    assert blocks["Прочие договоры#0"]["roles"]["executor"] == "Арендадатель BBC"


def test_kolonka_bez_shapki_i_nashi_yurlitsa_zhdut_cheloveka(finance_db):
    with finance_session() as session:
        workspace, batch = _start(session)
        assert set(batch.report["blocking"]) == {"columns", "entities"}
        entities = _section(batch, "entities")
        own = {tuple(item["names"]) for item in entities["items"] if item["own"]}
        # SAKOMPA-M и Sakompa-M — одно юрлицо; внешний продавец нашим не предложен.
        assert ("SAKOMPA-M", "Sakompa-M") in own
        assert all("Халык Актив" not in names for names in own)
        with pytest.raises(Exception):
            importer.apply(session, workspace, FULL, ACTOR, batch.id)


def _decide_all(session, workspace, batch, note_t=True):
    columns = {}
    for block in _section(batch, "columns")["items"]:
        for column in block["columns"]:
            if column["decision"].get("action") == "ask":
                columns[column["id"]] = {"action": "note" if note_t else "skip"}
    entities = _section(batch, "entities")
    own = [item["key"] for item in entities["items"] if item["own"]]
    return importer.decide(session, workspace, batch.id, {"columns": columns, "entities": {"confirmed": True, "own": own}})


def test_zavedenie_nichego_ne_teryaet(finance_db):
    with finance_session() as session:
        workspace, batch = _start(session)
        batch = _decide_all(session, workspace, batch)
        orphans = _section(batch, "orphans")
        assert {item["number"] for item in orphans["items"]} == {
            "Договор купли-продажи", "№041125", "№ 52-RENT"
        }
        diffs = _section(batch, "diffs")
        assert any(item["field"] == "amount" and item["main"] == "50 000" for item in diffs["items"])
        result = importer.apply(session, workspace, FULL, ACTOR, batch.id)
        assert result["created"] == 12 and not result["failed"]
        listing = service.list_all(session, workspace, FULL)
        registry = service.Registry(session, workspace)
        by_number = {}
        for item in listing["contracts"]:
            by_number.setdefault(item["values"]["number"], []).append(item)
        parties = listing["parties"]

        dsgroup = by_number["№041125"][0]["values"]
        assert parties[dsgroup["executor"]]["name"] == "ИП «DSGroup»"
        assert parties[dsgroup["customer"]]["own"] is True

        agent = by_number["№ AG-4"][0]
        assert agent["values"].get("amount") is None and agent["values"]["amount_terms"].startswith("20%")
        assert agent["values"]["billing"] == "terms"

        late = by_number["№ЮО/59"]
        assert {item["values"].get("planned_end_at") for item in late} >= {"2026-07-15"}
        assert all("number_taken" in {i["code"] for i in item["issues"]} for item in late)
        # Колонка T ушла в примечание, а не потерялась.
        assert any("у Бисултана" in (item["values"].get("note") or "") for item in late)

        unread = by_number["№ 51-RENT"][0]
        assert "unread_end_date" in {i["code"] for i in unread["issues"]}
        odd_status = agent["issues"]
        assert "status_unknown" in {i["code"] for i in odd_status}

        first = by_number["№ 17-BUH"][0]
        assert first["values"]["amendments_text"] == "№ 1 от 23.07.2024\n№2 10.12.2025"
        assert float(first["file_snapshot"]["paid"]) == 250000
        assert first["values"]["folder_url"].startswith("https://bitrix/")

        views = {view.title: view for view in registry.views}
        assert set(views) == {"Сводная", "Заказчик ГК", "Прочие договоры"}
        rent_members = [
            item for item in listing["contracts"]
            if any(v["view"] == views["Прочие договоры"].key for v in item["views"])
        ]
        assert {item["values"]["number"] for item in rent_members} == {"№ 50-RENT", "№ 51-RENT", "№ 52-RENT"}

        # Загруженный договор — история: перевод на другое ТОО спрашивает режим
        # даже у того, кто загрузил файл, в тот же день.
        owner = service.Actor(None, "owner@test")
        with pytest.raises(service.ModeRequired):
            service.patch(session, workspace, FULL, owner, __import__("uuid").UUID(first["id"]),
                          {"executor": "BBCA"}, known_seq=None)


def test_vygruzka_i_obratnaya_zagruzka(finance_db, tmp_path, monkeypatch):
    with finance_session() as session:
        workspace, batch = _start(session)
        batch = _decide_all(session, workspace, batch, note_t=False)
        importer.apply(session, workspace, FULL, ACTOR, batch.id)
        data = export.build(session, workspace, FULL, ACTOR)
        count = len(service.list_all(session, workspace, FULL)["contracts"])

    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(data))
    assert book.sheetnames == ["Сводная", "Заказчик ГК", "Прочие договоры"]
    buyer = book["Заказчик ГК"]
    headers = [cell.value for row in buyer.iter_rows(max_row=10) for cell in row if cell.value]
    # Шапка второго блока — как в файле: заказчик раньше исполнителя.
    assert headers.index("Заказчик") < headers.index("Исполнитель")

    # Выгрузка читается обратно тем же разбором и даёт те же договоры.
    from sqlalchemy import create_engine

    from app.finance import db as finance_db_module

    engine = create_engine(f"sqlite:///{tmp_path / 'again.db'}", future=True)
    monkeypatch.setattr(finance_db_module, "get_engine", lambda: engine)
    monkeypatch.setattr(finance_db_module, "_session_factory", None)
    monkeypatch.setattr(finance_db_module, "_initialized", False)
    with finance_session() as session:
        workspace = finance_service.ensure_workspace(session)
        again = importer.start(session, workspace, ACTOR, data, "выгрузка.xlsx")
        again = _decide_all(session, workspace, again, note_t=False)
        result = importer.apply(session, workspace, FULL, ACTOR, again.id)
        assert result["created"] == count
